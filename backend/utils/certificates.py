"""Fill an uploaded PowerPoint or Word certificate by replacing {{tags}} with project data."""
import calendar
import io
import re
from datetime import date

TAGS = {
    "client_name": "Client name",
    "project_name": "Project name",
    "location": "Project location",
    "service_text": "Work carried out",
    "reference": "Reference line",
    "completion_date": "Job completed date",
    "expiry_date": "Expiry date",
    "so_number": "SO number",
    "contract_period": "Contract period",
    "engineer_name": "Engineer name",
    "invoice_number": "Invoice number",
    "issue_date": "Date of issue",
}
TAG_PATTERN = re.compile(r"\{\{\s*([a-z_]+)\s*\}\}")
EXTENSIONS = {"pptx": b"ppt/", "docx": b"word/"}
DEFAULT_SERVICE_TEXT = "Domestic Water Tank and Pipeline Network."


def add_months(day, months):
    index = day.month - 1 + months
    year, month = day.year + index // 12, index % 12 + 1
    return date(year, month, min(day.day, calendar.monthrange(year, month)[1]))


def display_date(value):
    """DD/MM/YYYY, the format used on the certificate."""
    if isinstance(value, str):
        try:
            value = date.fromisoformat(value[:10])
        except ValueError:
            return value
    return value.strftime("%d/%m/%Y") if value else ""


def _fill_paragraph(paragraph, values):
    runs = paragraph.runs
    # Certificates carry no bold text; an explicit False also overrides bold inherited from the slide or style.
    for run in runs:
        run.font.bold = False
    text = "".join(run.text for run in runs)
    if "{{" not in text:
        return
    filled = TAG_PATTERN.sub(lambda match: str(values.get(match.group(1), match.group(0))), text)
    # Tags can be split across runs, so the first run keeps its formatting and takes the whole text.
    runs[0].text = filled
    for run in runs[1:]:
        run.text = ""


def _fill_pptx_shapes(shapes, values):
    for shape in shapes:
        if shape.shape_type == 6:  # group
            _fill_pptx_shapes(shape.shapes, values)
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                _fill_paragraph(paragraph, values)
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        _fill_paragraph(paragraph, values)


def _fill_docx_container(container, values):
    for paragraph in container.paragraphs:
        _fill_paragraph(paragraph, values)
    for table in getattr(container, "tables", []):
        for row in table.rows:
            for cell in row.cells:
                _fill_docx_container(cell, values)


def check_template(data, extension):
    """True when the bytes are a real Office file of the extension's kind that python can open."""
    import zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            if "[Content_Types].xml" not in names or not any(name.encode().startswith(EXTENSIONS[extension]) for name in names):
                return False
        template_tags(data, extension)
        return True
    except Exception:
        return False


def _text_of(data, extension):
    if extension == "pptx":
        from pptx import Presentation
        parts = []
        for slide in Presentation(io.BytesIO(data)).slides:
            stack = list(reversed(slide.shapes))
            while stack:
                shape = stack.pop()
                if shape.shape_type == 6:
                    stack.extend(reversed(shape.shapes))
                if shape.has_text_frame:
                    parts.extend(paragraph.text for paragraph in shape.text_frame.paragraphs)
                if getattr(shape, "has_table", False) and shape.has_table:
                    parts.extend(cell.text for row in shape.table.rows for cell in row.cells)
        return "\n".join(parts)
    from docx import Document
    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def template_tags(data, extension):
    """The tags found in the file, in first-seen order (unknown ones included so the page can warn about them)."""
    found = []
    for tag in TAG_PATTERN.findall(_text_of(data, extension)):
        if tag not in found:
            found.append(tag)
    return found


def fill_template(data, extension, values):
    output = io.BytesIO()
    if extension == "pptx":
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(data))
        for slide in presentation.slides:
            _fill_pptx_shapes(slide.shapes, values)
        presentation.save(output)
    else:
        from docx import Document
        document = Document(io.BytesIO(data))
        _fill_docx_container(document, values)
        for section in document.sections:
            for part in (section.header, section.footer):
                _fill_docx_container(part, values)
        document.save(output)
    return output.getvalue()
