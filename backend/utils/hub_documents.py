"""Files the Operations pages download: filtered tables as Excel or PDF, and reports as Word."""
import io
from datetime import date, datetime

from backend.utils.schedule_export import label_with_company


def safe_cell(value):
    """Text that a spreadsheet would run as a formula is stored as plain text."""
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def table_workbook(title, headers, rows, company=""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "".join(character for character in title if character not in "[]:*?/\\")[:31] or "Sheet1"
    label_with_company(workbook, sheet, company, title)
    sheet.append([str(header) for header in headers])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="142B35")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    widths = [len(str(header)) for header in headers]
    for row in rows:
        sheet.append([safe_cell(value) for value in row])
        for index, value in enumerate(row[:len(widths)]):
            widths[index] = max(widths[index], min(len(str(value or "")), 45))
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width + 3
    sheet.freeze_panes = "A2"
    raw = io.BytesIO()
    workbook.save(raw)
    raw.seek(0)
    return raw


def table_pdf(title, headers, rows, company=""):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from xml.sax.saxutils import escape

    raw = io.BytesIO()
    page = landscape(A4)
    margin = 10 * mm
    pdf = SimpleDocTemplate(raw, pagesize=page, title=title, leftMargin=margin, rightMargin=margin, topMargin=margin, bottomMargin=margin)
    styles = getSampleStyleSheet()
    size = 7 if len(headers) <= 12 else 5.5
    cell = ParagraphStyle("cell", parent=styles["Normal"], fontSize=size, leading=size + 1.5)
    head = ParagraphStyle("head", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    # Columns share the page width by how much text they hold, so a table of any width fits without running off the page.
    weights = [max(len(str(header)), 4) for header in headers]
    for row in rows[:400]:
        for index, value in enumerate(row[:len(weights)]):
            weights[index] = max(weights[index], min(len(str(value or "")), 40))
    available = page[0] - 2 * margin
    widths = [available * weight / sum(weights) for weight in weights] if weights else None
    data = [[Paragraph(escape(str(header)), head) for header in headers] or [""]]
    data += [[Paragraph(escape(str(value if value is not None else "")), cell) for value in row] for row in rows]
    table = Table(data, repeatRows=1, colWidths=widths)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#142B35")),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.lightgrey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
        ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]))
    heading = [Paragraph(escape(company), styles["Heading3"])] if company else []
    pdf.build([*heading, Paragraph(escape(title), styles["Heading4"]), Spacer(1, 4), table])
    raw.seek(0)
    return raw


MASTER_SHEET_COLUMNS = {
    "date": "date", "so no": "so", "dn n": "dn", "dn no": "dn", "dn": "dn", "customer": "customer_text",
    "delivery address": "address", "delivery person": "driver", "signed delivery note": "signed", "amount": "amount",
    "helper": "helpers",
}


def read_delivery_sheet(data):
    """The Delivery Master Sheet's columns as note fields; rows without a DN number are skipped."""
    from openpyxl import load_workbook

    sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip().lower() for value in next(rows)]
    if "dn n" not in headers and "dn no" not in headers and "dn" not in headers:
        raise ValueError("No DN column.")
    notes = []
    for row in rows:
        note = {}
        for header, value in zip(headers, row):
            field = MASTER_SHEET_COLUMNS.get(header)
            if field and value not in (None, ""):
                if isinstance(value, (datetime, date)):
                    value = value.strftime("%Y-%m-%d")
                note[field] = value
        if note.get("dn"):
            notes.append(note)
    return notes


def report_docx(report, image_bytes, company_images, company=""):
    """A saved report as Word: cover, document control, letter, contents-order blocks, and the footer bar."""
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Cm, Mm, Pt, RGBColor

    doc = report.get("doc") or {}
    word = Document()
    section = word.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Cm(2)
    style = word.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    def picture(paragraph, file_id, width):
        data = image_bytes(file_id) if file_id else None
        if data:
            try:
                paragraph.add_run().add_picture(io.BytesIO(data), width=width)
                return True
            except Exception:
                paragraph.add_run("[This image could not be added]")
        return False

    def text(value, bold=False, size=None, center=False, color=None):
        paragraph = word.add_paragraph()
        run = paragraph.add_run(str(value or ""))
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)
        if center:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        return paragraph

    def table(rows, widths=None):
        rows = [row for row in rows if row]
        if not rows:
            return
        columns = max(len(row) for row in rows)
        grid = word.add_table(rows=len(rows), cols=columns)
        grid.style = "Table Grid"
        for row_index, row in enumerate(rows):
            for column, value in enumerate(row):
                grid.cell(row_index, column).text = str(value if value is not None else "")
        word.add_paragraph()

    if section.different_first_page_header_footer is False:
        section.different_first_page_header_footer = True
    if company_images.get("footer"):
        picture(section.footer.paragraphs[0], company_images["footer"], Cm(17))
        picture(section.first_page_footer.paragraphs[0], company_images["footer"], Cm(17))

    if report.get("kind") == "Certificate":
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = Mm(297), Mm(210)
        fields = doc.get("certificate") or {}
        if company_images.get("logo"):
            picture(word.add_paragraph(), company_images["logo"], Cm(3))
        text(fields.get("company") or company, bold=True, size=16, center=True)
        text(fields.get("heading") or "Certificate", bold=True, size=26, center=True, color="147D68")
        text(fields.get("wording"), center=True)
        table([
            ["Client", fields.get("client")], ["Project / Villa", fields.get("project")], ["Location", fields.get("location")],
            ["Work Description", fields.get("work")], ["Reference", fields.get("reference")], ["Engineer", fields.get("engineer")],
            ["Job Completed", fields.get("completed")], ["Expiry", fields.get("expiry")],
        ])
        if company_images.get("stamp"):
            picture(word.add_paragraph(), company_images["stamp"], Cm(4))
    else:
        cover = doc.get("cover") or {}
        if company_images.get("logo"):
            logo = word.add_paragraph()
            logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture(logo, company_images["logo"], Cm(4))
        text(cover.get("title") or report.get("title"), bold=True, size=26, center=True, color="142B35")
        text(cover.get("subtitle"), size=15, center=True)
        for label, name in ((cover.get("label"), cover.get("name")), (cover.get("label2"), cover.get("name2"))):
            if label or name:
                text(label, bold=True, center=True)
                text(name, center=True)
        word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        lines = doc.get("header") or {}
        parts = [part for part in (lines.get("document"), lines.get("project"), lines.get("client")) if part]
        if parts:
            header_table = section.header.add_table(1, len(parts), Cm(17))
            header_table.style = "Table Grid"
            for index, part in enumerate(parts):
                header_table.cell(0, index).text = str(part)
        control = doc.get("control") or {}
        if control.get("enabled"):
            text(control.get("title1") or "Document Control", bold=True, size=14)
            table([[row.get("label"), row.get("value"), row.get("label2"), row.get("value2")] for row in control.get("rows") or []])
            if control.get("title2"):
                text(control["title2"], bold=True, size=12)
            if control.get("prepared") or control.get("approved"):
                table([["Prepared By", control.get("prepared")], ["Approved By", control.get("approved")]])
            word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        letter = doc.get("letter") or {}
        if letter.get("enabled"):
            for label, field in (("From", "from"), ("To", "to"), ("Project", "project"), ("Subject", "subject")):
                if letter.get(field):
                    text(f"{label}: {letter[field]}")
            text(letter.get("body"))
            for item in letter.get("includes") or []:
                word.add_paragraph(str(item), style="List Bullet")
            text(letter.get("signer"), bold=True)
            word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        current_page, heading = 1, 0
        for block in doc.get("blocks") or []:
            page = int(block.get("page") or current_page)
            if page != current_page:
                word.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
                current_page = page
            kind = block.get("type")
            if kind == "heading":
                heading += 1
                word.add_heading(f"{heading}. {block.get('text') or ''}", level=1)
            elif kind == "subheading":
                word.add_heading(block.get("text") or "", level=2)
            elif kind == "text":
                text(block.get("text"))
            elif kind == "list":
                for item in block.get("items") or []:
                    word.add_paragraph(str(item), style="List Number" if block.get("numbered") else "List Bullet")
            elif kind == "table":
                table(block.get("rows") or [])
            elif kind == "keyvalue":
                table([[row.get("key"), row.get("value")] for row in block.get("rows") or []])
            elif kind == "photos":
                per_row = max(1, min(int(block.get("columns") or 2), 6))
                photos = block.get("photos") or []
                width = Cm(17 * (float(block.get("size") or 100) / 100) / per_row - 0.2)
                for start in range(0, len(photos), per_row):
                    line = word.add_paragraph()
                    for photo in photos[start:start + per_row]:
                        picture(line, photo.get("file_id"), width)
    double_page_border(section)
    raw = io.BytesIO()
    word.save(raw)
    raw.seek(0)
    return raw


def double_page_border(section):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    borders = OxmlElement("w:pgBorders")
    borders.set(qn("w:offsetFrom"), "page")
    for side in ("top", "left", "bottom", "right"):
        border = OxmlElement(f"w:{side}")
        for name, value in (("val", "double"), ("sz", "6"), ("space", "24"), ("color", "147D68")):
            border.set(qn(f"w:{name}"), value)
        borders.append(border)
    section._sectPr.append(borders)
