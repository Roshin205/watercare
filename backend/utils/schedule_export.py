"""Client schedule layout for the AMC contracts export: one block per client, its sites listed beneath it."""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADERS = ["SL NO", "Client", "", "Project", "SO Number", "Contract Period", "System Details",
           "Engineer Visit", "Microbiology", "Legionella"]
COLUMN_WIDTHS = (8, 24, 5, 28, 12, 32, 44, 16, 16, 16)


def month_year(value):
    try:
        return date.fromisoformat(str(value)[:10]).strftime("%B %Y")
    except ValueError:
        return ""


def schedule_rows(contracts):
    """One list per contract. SL NO and the client name appear on the first row of each client only."""
    groups = {}
    for row in contracts:
        groups.setdefault(row.get("client") or "No client", []).append(row)
    table = []
    for serial, name in enumerate(sorted(groups, key=str.lower), start=1):
        for number, row in enumerate(groups[name], start=1):
            start, end = month_year(row.get("start")), month_year(row.get("expiry"))
            period = row.get("period_text") or (f"{start} to {end}" if start and end and end != start else start or end)
            table.append([
                serial if number == 1 else "", name if number == 1 else "", number, row.get("site") or "", row.get("so") or "",
                period, row.get("system_details") or "", row.get("frequency") or "",
                "Yes" if row.get("microbiology") == "Yes" else "NO", "Yes" if row.get("legionella") == "Yes" else "NO",
            ])
    return table


def schedule_workbook(table, clean, company=""):
    """Build the sheet; `clean` neutralises spreadsheet formulas in user text. The company name goes in the printed
    page header and the file's properties, so the grid keeps the client's layout."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Client schedule"
    label_with_company(workbook, sheet, company, "Client schedule")
    thin = Side(style="thin", color="000000")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    green, yellow = PatternFill("solid", fgColor="92D86B"), PatternFill("solid", fgColor="FFFF00")
    sheet.append(HEADERS)
    for cell in sheet[1]:
        cell.font, cell.fill, cell.border = Font(bold=True), green, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 22
    block_start = 2
    for index, values in enumerate(table, start=2):
        if values[1] and index > 2:
            _merge_client(sheet, block_start, index - 1)
            block_start = index
        sheet.append([clean(value) for value in values])
        for column, cell in enumerate(sheet[index], start=1):
            cell.border = border
            cell.alignment = Alignment(
                horizontal="center" if column in (1, 3, 5) else "left", vertical="center", wrap_text=True,
            )
            if column == 3:
                cell.font = Font(bold=True)
        sheet.cell(index, 2).fill = yellow
        sheet.cell(index, 2).font = Font(bold=True)
        sheet.row_dimensions[index].height = max(20, 15 * (str(values[6]).count("\n") + 1))
    if table:
        _merge_client(sheet, block_start, len(table) + 1)
    for letter, width in zip("ABCDEFGHIJ", COLUMN_WIDTHS):
        sheet.column_dimensions[letter].width = width
    sheet.freeze_panes = "A2"
    return workbook


def label_with_company(workbook, sheet, company, title):
    if company:
        workbook.properties.creator = company
        workbook.properties.title = f"{company} - {title}"
        sheet.oddHeader.center.text = company.replace("&", "&&")  # "&" starts a header code in Excel
        sheet.oddHeader.center.size = 12


def _merge_client(sheet, first, last):
    if last > first:
        sheet.merge_cells(start_row=first, start_column=1, end_row=last, end_column=1)
        sheet.merge_cells(start_row=first, start_column=2, end_row=last, end_column=2)
