from datetime import date, datetime, timezone
from io import BytesIO
from uuid import uuid4

import pytest
from werkzeug.security import generate_password_hash

from backend.database import db
from backend.models import memberships, users
from backend.utils import hub_logic as logic
from tests.test_app import ADMIN_KEY, app, client, csrf  # noqa: F401 - shared fixtures


def post(client, path, body=None):
    return client.post(f"/api/hub{path}", json=body or {}, headers=csrf(client))


def patch(client, path, body):
    return client.patch(f"/api/hub{path}", json=body, headers=csrf(client))


def upload(client, path, name, data, **fields):
    return client.post(f"/api/hub{path}", data={"file": (BytesIO(data), name), **fields},
                       headers=csrf(client), content_type="multipart/form-data")


def contract(client, **fields):
    body = {"client": "Daikin", "site": "True Bell", "so": "1087", "start": "2026-07-01", "expiry": "2027-06-30",
            "frequency": "Quarterly", **fields}
    response = post(client, "/contracts", body)
    assert response.status_code == 201, response.json
    return response.json["data"]


def role_client(app, role, name="Viewer"):
    now = datetime.now(timezone.utc).isoformat()
    user_id = str(uuid4())
    db.session.execute(users.insert().values(
        id=user_id, email=f"{role}@example.test", full_name=name,
        password_hash=generate_password_hash("secure-test-password"),
        access_key_hash=generate_password_hash(ADMIN_KEY), is_active=True, created_at=now, updated_at=now,
    ))
    db.session.execute(memberships.insert().values(
        id=str(uuid4()), organization_id="org-1", user_id=user_id, role=role, status="active", created_at=now,
    ))
    db.session.commit()
    test_client = app.test_client()
    admin_level = role in ("service_coordinator", "accountant")
    login = {"email": f"{role}@example.test", "password": "secure-test-password", "login_type": "admin" if admin_level else "staff"}
    if admin_level:
        login["access_key"] = ADMIN_KEY
    response = test_client.post("/api/auth/login", json=login)
    assert response.status_code == 200, response.json
    test_client.csrf_token = response.json["csrf_token"]
    return test_client


def test_schedule_fills_due_months_by_frequency_inside_the_contract():
    schedule = logic.fill_schedule({"frequency": "Quarterly", "start": "2026-07-01", "expiry": "2027-06-30"}, date(2026, 10, 30))
    assert sorted(schedule) == ["2026-10", "2027-01", "2027-04"]
    assert all(visit["status"] == "due" for visit in schedule.values())


def test_schedule_keeps_done_missed_and_na_months_and_leaves_on_call_alone():
    today = date(2026, 10, 30)
    kept = {"2026-10": {"status": "done", "date": "2026-10-05", "engineer": "A"}, "2027-02": {"status": "na", "date": "2027-02-01"}}
    schedule = logic.fill_schedule({"frequency": "Bimonthly", "start": "2026-10-01", "expiry": "2027-04-30", "schedule": kept}, today)
    assert schedule["2026-10"]["status"] == "done" and schedule["2027-02"]["status"] == "na"
    assert schedule["2026-12"]["status"] == "due" and "2026-11" not in schedule
    assert logic.fill_schedule({"frequency": "On call", "start": "2026-10-01", "expiry": "2027-04-30", "schedule": kept}, today) == kept


def test_changing_frequency_removes_future_due_months_it_no_longer_wants():
    today = date(2026, 10, 30)
    row = {"frequency": "Monthly", "start": "2026-10-01", "expiry": "2026-12-31"}
    row["schedule"] = logic.fill_schedule(row, today)
    row["frequency"] = "Quarterly"
    assert sorted(logic.fill_schedule(row, today)) == ["2026-10"]


def test_contract_gets_defaults_a_client_and_next_due(client):
    row = contract(client)
    assert row["status"] == "Active" and row["report"] == "Pending" and row["invoice"] == "Pending"
    assert row["client_id"] and row["next_due"] == min(row["schedule"]) + "-01"
    assert [item["name"] for item in client.get("/api/clients").json["data"]] == ["Daikin"]


def test_log_visit_marks_month_done_and_moves_the_dates(client):
    row = contract(client)
    month = min(row["schedule"])
    saved = post(client, f"/contracts/{row['id']}/visit", {
        "date": f"{month}-12", "engineer": "Worker 17", "report": "Submitted", "lab": "Received",
    }).json["data"]
    assert saved["schedule"][month] == {"status": "done", "date": f"{month}-12", "engineer": "Worker 17"}
    assert saved["last_visit"] == f"{month}-12" and saved["visited_by"] == "Worker 17"
    assert saved["report"] == "Submitted" and saved["lab"] == "Received"
    assert saved["next_due"] > saved["last_visit"]


def test_schedule_cell_cycle_move_and_remove(client):
    row = contract(client)
    month = min(row["schedule"])
    saved = post(client, f"/contracts/{row['id']}/schedule", {"month": month, "status": "missed"}).json["data"]
    assert saved["schedule"][month]["status"] == "missed"
    moved = post(client, f"/contracts/{row['id']}/schedule", {"month": month, "move_to": "2030-05"}).json["data"]
    assert month not in moved["schedule"] and moved["schedule"]["2030-05"]["date"] == "2030-05-01"
    gone = post(client, f"/contracts/{row['id']}/schedule", {"month": "2030-05", "remove": True}).json["data"]
    assert "2030-05" not in gone["schedule"]
    assert post(client, f"/contracts/{row['id']}/schedule", {"month": "nope"}).status_code == 400


def test_invoice_submitted_fills_today(client):
    row = contract(client)
    assert patch(client, f"/contracts/{row['id']}", {"invoice": "Submitted"}).json["data"]["invoice_date"] == date.today().isoformat()


def test_close_reopen_and_renew_with_copies(client):
    row = contract(client)
    assert post(client, f"/contracts/{row['id']}/close", {"close_reason": "Terminated"}).json["data"]["status"] == "Closed"
    assert post(client, f"/contracts/{row['id']}/reopen").json["data"]["status"] == "Active"
    post(client, "/samples", {"contract_id": row["id"], "site": row["site"], "test": "Legionella", "duration": "Quarterly", "collection_date": "2026-10-01"})
    post(client, "/ppm_planners", {"contract_id": row["id"], "site": row["site"], "year": 2026, "groups": []})
    post(client, "/chemical_lines", {"contract_id": row["id"], "site": row["site"], "chemical": "SANITER 420", "contract_qty": 2, "months": {"2026-06": 2}})
    renewed = post(client, f"/contracts/{row['id']}/renew", {
        "so": "2000", "period_text": "Jul 27 to Jun 28", "start": "2027-07-01", "expiry": "2028-06-30",
        "copy_samples": True, "copy_ppm": True, "copy_chemicals": True,
    })
    assert renewed.status_code == 201
    new = renewed.json["data"]
    assert new["so"] == "2000" and new["renewed_from"] == row["id"]
    rows = {item["id"]: item for item in client.get("/api/hub/contracts").json["data"]}
    assert rows[row["id"]]["status"] == "Closed" and rows[row["id"]]["close_reason"] == "Renewed"
    assert rows[row["id"]]["renewed_to"] == new["id"]
    samples = client.get(f"/api/hub/samples?parent={new['id']}").json["data"]
    assert len(samples) == 1 and samples[0]["status"] == "Pending" and not samples[0]["collection_date"]
    assert len(client.get(f"/api/hub/ppm_planners?parent={new['id']}").json["data"]) == 1
    lines = client.get(f"/api/hub/chemical_lines?parent={new['id']}").json["data"]
    assert len(lines) == 1 and lines[0]["months"] == {} and lines[0]["so"] == "2000"


def test_grn_status_marks_matching_contract_pending(client):
    contract(client)
    post(client, "/grn_projects", {"project": "Other", "so": "1087", "status": "Waiting GRN"})
    assert client.get("/api/hub/contracts").json["data"][0]["grn"] == "Pending"


def test_sample_dates_and_collect(client):
    created = post(client, "/samples", {"site": "True Bell", "test": "Legionella", "duration": "Quarterly", "collection_date": "2026-10-18"}).json["data"]
    assert created["next_due"] == "2027-01-16" and created["expiry"] == "2027-01-16"
    once = post(client, "/samples", {"site": "True Bell", "test": "TPC", "duration": "Once", "collection_date": "2026-10-18"}).json["data"]
    assert not once["next_due"]
    manual = post(client, "/samples", {"site": "True Bell", "test": "TPC", "duration": "Quarterly", "collection_date": "2026-10-18", "next_due": "2026-11-18"}).json["data"]
    assert manual["next_due"] == "2026-11-18"
    collected = post(client, f"/samples/{created['id']}/collect", {"collection_date": "2026-11-01", "collected_by": "Ali"}).json["data"]
    assert collected["status"] == "Collected" and collected["report"] == "Not submitted"
    assert collected["next_due"] == "2027-01-30" and collected["expiry"] == "2027-01-30"


def test_generate_planners_uses_samples_or_the_default(client):
    with_samples = contract(client, site="A")
    contract(client, site="B")
    post(client, "/samples", {"contract_id": with_samples["id"], "site": "A", "test": "Legionella", "duration": "Quarterly"})
    assert post(client, "/ppm_planners/generate", {"year": 2027}).json["created"] == 2
    assert post(client, "/ppm_planners/generate", {"year": 2027}).json["created"] == 0
    planners = {item["site"]: item for item in client.get("/api/hub/ppm_planners").json["data"]}
    assert [row["service"] for row in planners["A"]["groups"][0]["rows"]] == ["Legionella Analysis"]
    assert planners["A"]["groups"][0]["rows"][0]["months"] == [1, "", "", 1, "", "", 1, "", "", 1, "", ""]
    assert len(planners["B"]["groups"][0]["rows"]) == 3


def test_chemical_delivery_moves_stock_and_clearing_restores_it(client):
    item = post(client, "/inventory_items", {"name": "SANITER 420", "opening_stock": 10, "unit_cost": 5}).json["data"]
    line = post(client, "/chemical_lines", {"site": "True Bell", "chemical": "saniter 420", "contract_qty": 4}).json["data"]
    assert post(client, f"/chemical_lines/{line['id']}/month", {"month": "2026-06", "qty": 3}).json["message"] == ""
    moves = client.get("/api/hub/inventory_moves").json["data"]
    assert len(moves) == 1 and moves[0]["type"] == "OUT" and moves[0]["qty"] == 3 and moves[0]["item_id"] == item["id"]
    post(client, f"/chemical_lines/{line['id']}/month", {"month": "2026-06", "qty": 5})
    assert [move["qty"] for move in client.get("/api/hub/inventory_moves").json["data"]] == [5]
    assert client.get("/api/hub/overview").json["inventory"]["value"] == 25
    post(client, f"/chemical_lines/{line['id']}/month", {"month": "2026-06", "qty": ""})
    assert client.get("/api/hub/inventory_moves").json["data"] == []
    assert client.get("/api/hub/chemical_lines").json["data"][0]["months"] == {}


def test_chemical_delivery_without_a_matching_item_says_so_and_can_be_switched_off_or_capped(client):
    line = post(client, "/chemical_lines", {"site": "X", "chemical": "Unknown", "contract_qty": 4}).json["data"]
    assert "stock was not changed" in post(client, f"/chemical_lines/{line['id']}/month", {"month": "2026-06", "qty": 1}).json["message"]
    client.put("/api/hub/settings", json={"inventory_link": False, "cap_deliveries": True}, headers=csrf(client))
    capped = post(client, f"/chemical_lines/{line['id']}/month", {"month": "2026-07", "qty": 9}).json
    assert capped["data"]["months"]["2026-07"] == 3 and "Limited" in capped["message"]


def test_delivery_note_numbers_duplicates_and_tracker(client):
    year = str(date.today().year)[2:]
    first = post(client, "/delivery_notes", {
        "date": "2026-10-26", "client": "Daikin", "so": "1087", "site": "True Bell",
        "items": [{"name": "SANITER 420", "qty": 2, "unit": "Drum"}],
    }).json["data"]
    assert first["dn"] == f"DN-{year}-001" and first["customer_text"] == f"Delivery Note – DN-{year}-001 – Daikin – 1087"
    assert post(client, "/delivery_notes", {"date": "2026-10-26"}).json["data"]["dn"] == f"DN-{year}-002"
    assert post(client, "/delivery_notes", {"date": "2026-10-26", "dn": first["dn"]}).status_code == 409
    client.put("/api/hub/settings", json={"dn_start": 514}, headers=csrf(client))
    assert post(client, "/delivery_notes", {"date": "2026-10-26"}).json["data"]["dn"] == f"DN-{year}-514"
    assert post(client, f"/delivery_notes/{first['id']}/to-chemicals").json["added"] == 1
    lines = client.get("/api/hub/chemical_lines").json["data"]
    assert lines[0]["chemical"] == "SANITER 420" and lines[0]["months"] == {"2026-10": 2}


def test_master_sheet_import_adds_and_updates_by_dn(client):
    from openpyxl import Workbook
    book = Workbook()
    book.active.append(["S.No", "Date", "SO No", "DN N", "Customer", "Delivery Address", "Delivery person", "Signed Delivery Note", "Amount", "Helper"])
    book.active.append([1, date(2026, 10, 1), "770", "DN-26-900", "Text", "Dubai", "Sam", "Collected", 500, "Ali"])
    raw = BytesIO()
    book.save(raw)
    assert upload(client, "/delivery_notes/import", "sheet.xlsx", raw.getvalue()).json == {"added": 1, "updated": 0}
    assert upload(client, "/delivery_notes/import", "sheet.xlsx", raw.getvalue()).json == {"added": 0, "updated": 1}
    note = client.get("/api/hub/delivery_notes").json["data"][0]
    assert note["driver"] == "Sam" and note["amount"] == 500 and note["date"] == "2026-10-01"


def test_staff_csv_adds_and_updates(client):
    data = b"Name,Position,Employee ID\nCanutte,Technician,E1\n"
    assert upload(client, "/staff/import", "s.csv", data).json == {"added": 1, "updated": 0}
    assert upload(client, "/staff/import", "s.csv", b"Name,Position,Employee ID\nNew Name,Lead,E1\n").json == {"added": 0, "updated": 1}
    assert client.get("/api/hub/staff").json["data"][0]["position"] == "Lead"


def test_attendance_marks_are_separate_and_clearing_removes_one_day(client):
    person = post(client, "/staff", {"name": "Canutte"}).json["data"]
    marks = [{"staff_id": person["id"], "date": "2026-10-01", "code": "P"}, {"staff_id": person["id"], "date": "2026-10-02", "code": "A"}]
    assert client.put("/api/hub/attendance", json={"marks": marks}, headers=csrf(client)).status_code == 200
    client.put("/api/hub/attendance", json={"marks": [{"staff_id": person["id"], "date": "2026-10-01", "code": ""}]}, headers=csrf(client))
    rows = client.get("/api/hub/attendance?period=2026-10").json["data"]
    assert [(row["date"], row["code"]) for row in rows] == [("2026-10-02", "A")]


def test_sections_are_seeded_once_calibration_due_is_automatic_and_custom_sections_work(client):
    systems = {section["system"]: section for section in client.get("/api/hub/bootstrap").json["sections"]}
    assert set(systems) == {"assets", "calibration"}
    assert client.delete(f"/api/hub/sections/{systems['assets']['id']}", headers=csrf(client)).status_code == 400
    row = post(client, "/section_rows", {"section_id": systems["calibration"]["id"], "values": {"instrument": "Meter", "calibration_date": "2026-01-31"}}).json["data"]
    assert row["values"]["calibration_due"] == "2027-01-31"
    assert post(client, "/section_rows", {"section_id": systems["calibration"]["id"], "values": {"model": "No name"}}).status_code == 400
    custom = post(client, "/sections", {"name": "Vehicles", "columns": [{"key": "plate", "label": "Plate", "type": "text"}]}).json["data"]
    assert post(client, "/sections", {"name": "Bad", "columns": []}).status_code == 400
    assert post(client, "/section_rows", {"section_id": custom["id"], "values": {"plate": "D 123"}}).status_code == 201
    assert post(client, "/section_rows", {"section_id": custom["id"], "values": {"other": "x"}}).status_code == 400
    assert client.delete(f"/api/hub/sections/{custom['id']}", headers=csrf(client)).status_code == 200
    assert client.get(f"/api/hub/section_rows?parent={custom['id']}").json["data"] == []
    assert len(client.get("/api/hub/bootstrap").json["sections"]) == 2


def test_validation_refuses_unknown_fields_and_bad_dates(client):
    assert post(client, "/contracts", {"site": "A", "owner": "x"}).status_code == 400
    assert post(client, "/contracts", {"client": "x"}).status_code == 400
    assert post(client, "/contracts", {"site": "A", "start": "01-01-2026"}).status_code == 400
    assert post(client, "/nothing", {}).status_code == 404


def test_files_upload_list_download_delete(client):
    person = post(client, "/staff", {"name": "A"}).json["data"]
    assert upload(client, "/files", "p.png", b"\x89PNG\r\n\x1a\nabc", record_id=person["id"], kind="Photo").status_code == 201
    listed = client.get(f"/api/hub/files?record_id={person['id']}").json["data"]
    assert listed[0]["image_type"] == "Photo"
    assert client.get(f"/api/hub/files/{listed[0]['id']}").status_code == 200
    assert upload(client, "/files", "p.png", b"nope", record_id=person["id"]).status_code == 400
    assert client.delete(f"/api/hub/staff/{person['id']}", headers=csrf(client)).status_code == 200
    assert client.get(f"/api/hub/files?record_id={person['id']}").json["data"] == []


def test_file_index_gives_the_newest_photo_and_hides_identity_kinds(app, client):
    person = post(client, "/staff", {"name": "A"}).json["data"]
    upload(client, "/files", "old.png", b"\x89PNG\r\n\x1a\nold", record_id=person["id"], kind="Photo")
    newest = upload(client, "/files", "new.png", b"\x89PNG\r\n\x1a\nnew", record_id=person["id"], kind="Photo").json["data"]
    upload(client, "/files", "p.pdf", b"%PDF-1.4", record_id=person["id"], kind="Passport")
    entry = client.get("/api/hub/file-index?collection=staff").json["data"][person["id"]]
    assert entry["photo"] == newest["id"] and sorted(entry["kinds"]) == ["Passport", "Photo"]
    viewer = role_client(app, "viewer").get("/api/hub/file-index?collection=staff").json["data"][person["id"]]
    assert viewer["kinds"] == ["Photo"]


@pytest.mark.parametrize("output,signature", [("xlsx", b"PK"), ("pdf", b"%PDF")])
def test_export_returns_the_rows_sent(client, output, signature):
    response = post(client, "/export", {"title": "AMC Contracts", "format": output, "headers": ["A", "B"], "rows": [["=1+1", 2]]})
    assert response.status_code == 200 and response.data.startswith(signature)


def test_report_word_export(client):
    report = {"kind": "Tank cleaning report", "title": "T", "doc": {
        "cover": {"title": "Report"}, "control": {"enabled": True, "rows": [{"label": "A", "value": "B"}]},
        "letter": {"enabled": True, "body": "Hello", "includes": ["One"]},
        "blocks": [{"type": "heading", "text": "H", "page": 1}, {"type": "table", "rows": [["a", "b"]], "page": 2},
                   {"type": "list", "items": ["x"], "page": 2}, {"type": "keyvalue", "rows": [{"key": "k", "value": "v"}], "page": 2}],
    }}
    assert post(client, "/reports/word", {"report": report}).data.startswith(b"PK")
    certificate = {"kind": "Certificate", "doc": {"certificate": {"heading": "Certificate", "client": "C"}}}
    assert post(client, "/reports/word", {"report": certificate}).data.startswith(b"PK")


def test_viewer_reads_but_cannot_change(app, client):
    row = contract(client)
    viewer = role_client(app, "viewer")
    assert len(viewer.get("/api/hub/contracts").json["data"]) == 1
    assert post(viewer, "/contracts", {"site": "X"}).status_code == 403
    assert patch(viewer, f"/contracts/{row['id']}", {"so": "1"}).status_code == 403
    assert post(viewer, f"/contracts/{row['id']}/close").status_code == 403


def test_employee_sees_only_their_visits_and_no_other_section(app, client):
    row = contract(client)
    month, other_month = min(row["schedule"]), max(row["schedule"])
    post(client, f"/contracts/{row['id']}/schedule", {"month": month, "engineer": "Sam"})
    post(client, f"/contracts/{row['id']}/schedule", {"month": other_month, "engineer": "Ali"})
    contract(client, site="Not mine")
    employee = role_client(app, "employee", "Sam")
    rows = employee.get("/api/hub/contracts").json["data"]
    assert len(rows) == 1 and list(rows[0]["schedule"]) == [month] and "invoice" not in rows[0]
    assert employee.get("/api/hub/staff").status_code == 403
    assert employee.get("/api/clients").status_code == 403
    assert employee.get("/api/hub/overview").status_code == 403


def test_identity_details_are_limited_to_admin_and_accountant(app, client):
    person = post(client, "/staff", {"name": "A", "passport_no": "P1", "passport_expiry": "2026-11-20"}).json["data"]
    stored = upload(client, "/files", "p.pdf", b"%PDF-1.4", record_id=person["id"], kind="Passport").json["data"]
    coordinator = role_client(app, "service_coordinator", "Coordinator")
    row = coordinator.get("/api/hub/staff").json["data"][0]
    assert "passport_no" not in row and row["name"] == "A"
    assert patch(coordinator, f"/staff/{person['id']}", {"passport_no": "X"}).status_code == 403
    assert patch(coordinator, f"/staff/{person['id']}", {"position": "Tech"}).status_code == 200
    assert coordinator.get(f"/api/hub/files?record_id={person['id']}").json["data"] == []
    assert coordinator.get(f"/api/hub/files/{stored['id']}").status_code == 403
    assert coordinator.get("/api/hub/overview").json["staff"]["documents"] == []
    assert client.get("/api/hub/staff").json["data"][0]["passport_no"] == "P1"
    assert role_client(app, "accountant", "Accounts").get("/api/hub/staff").json["data"][0]["passport_no"] == "P1"


def test_calendar_blends_visits_and_sample_due_dates(app, client):
    row = contract(client, latitude=25.2, longitude=55.3)
    month = min(row["schedule"])
    post(client, f"/contracts/{row['id']}/schedule", {"month": month, "engineer": "Sam"})
    post(client, "/samples", {"contract_id": row["id"], "site": row["site"], "test": "Legionella", "next_due": f"{month}-20"})
    first, last = f"{month}-01", f"{month}-28"
    items = client.get(f"/api/hub/calendar?from={first}&to={last}").json["data"]
    assert {item["record_type"] for item in items} == {"visit", "sample"}
    assert client.get(f"/api/hub/calendar?from={first}&to={last}&client_id=other").json["data"] == []
    assert client.get("/api/hub/calendar").status_code == 400
    located = client.get("/api/hub/site-locations").json["data"]
    assert located[0]["site"] == "True Bell" and located[0]["latitude"] == 25.2
    employee = role_client(app, "employee", "Sam")
    own = employee.get(f"/api/hub/calendar?from={first}&to={last}").json["data"]
    assert [item["record_type"] for item in own] == ["visit"]


PHOTO = b"\x89PNG\r\n\x1a\nphoto"


def visit_for(client, engineer="Sam"):
    row = contract(client, engineer=engineer)
    return row, min(row["schedule"])


def complete(test_client, row, month, **fields):
    return post(test_client, f"/contracts/{row['id']}/complete", {"month": month, "date": date.today().isoformat(), **fields})


def add_photos(test_client, row, month):
    for side in ("Before", "After"):
        assert upload(test_client, "/files", f"{side}.png", PHOTO, record_id=row["id"], kind=f"{side}:{month}").status_code == 201


def test_assigned_engineer_goes_on_each_new_visit(client):
    row, month = visit_for(client)
    assert row["schedule"][month]["engineer"] == "Sam"


def test_visit_completion_needs_photos_then_approval_and_keeps_a_history(app, client):
    row, month = visit_for(client)
    employee = role_client(app, "employee", "Sam")
    assert complete(employee, row, month).status_code == 400
    add_photos(employee, row, month)
    submitted = complete(employee, row, month, notes="All clean")
    assert submitted.status_code == 200 and submitted.json["data"]["approval"] == "Pending Approval"
    assert submitted.json["data"]["completion"]["notes"] == "All clean"
    assert complete(employee, row, month).status_code == 409

    waiting = client.get("/api/hub/approvals").json["pending"]
    assert len(waiting) == 1 and waiting[0]["site"] == "True Bell" and waiting[0]["month"] == month
    assert employee.get("/api/hub/approvals").status_code == 403
    assert post(employee, f"/contracts/{row['id']}/approve", {"month": month}).status_code == 403
    assert post(client, f"/contracts/{row['id']}/reject", {"month": month}).status_code == 400
    rejected = post(client, f"/contracts/{row['id']}/reject", {"month": month, "reason": "Photos are blurred"}).json["data"]
    assert rejected["schedule"][month]["approval"] == "Rejected" and rejected["schedule"][month]["status"] == "due"

    assert complete(employee, row, month).status_code == 200
    approved = post(client, f"/contracts/{row['id']}/approve", {"month": month}).json["data"]
    visit = approved["schedule"][month]
    assert visit["status"] == "done" and visit["approval"] == "Approved" and approved["visited_by"] == "Sam"
    assert approved["last_visit"] == date.today().isoformat()
    assert complete(employee, row, month).status_code == 409
    actions = [entry["action"] for entry in client.get("/api/hub/approvals").json["history"]]
    assert sorted(actions) == ["Approved", "Rejected", "Resubmitted", "Submitted"]


def test_staff_reach_only_their_own_completion_photos(app, client):
    mine, month = visit_for(client)
    other, other_month = visit_for(client, engineer="Ali")
    employee = role_client(app, "employee", "Sam")
    assert upload(employee, "/files", "b.png", PHOTO, record_id=other["id"], kind=f"Before:{other_month}").status_code == 403
    assert upload(employee, "/files", "b.png", PHOTO, record_id=mine["id"], kind="Contract").status_code == 403
    add_photos(employee, mine, month)
    add_photos(client, other, other_month)
    assert len(employee.get(f"/api/hub/files?record_id={mine['id']}").json["data"]) == 2
    assert employee.get(f"/api/hub/files?record_id={other['id']}").json["data"] == []
    theirs = client.get(f"/api/hub/files?record_id={other['id']}").json["data"][0]
    assert employee.get(f"/api/hub/files/{theirs['id']}").status_code == 403
    assert complete(employee, other, other_month).status_code == 403


def test_office_log_visit_needs_no_approval(client):
    row, month = visit_for(client)
    saved = post(client, f"/contracts/{row['id']}/visit", {"date": f"{month}-05", "engineer": "Sam"}).json["data"]
    assert saved["schedule"][month]["status"] == "done" and "approval" not in saved["schedule"][month]


def test_visit_cannot_move_onto_a_month_that_already_has_one(client):
    row = contract(client, frequency="Monthly")
    first, second = sorted(row["schedule"])[:2]
    assert post(client, f"/contracts/{row['id']}/schedule", {"month": first, "move_to": second}).status_code == 400
    moved = post(client, f"/contracts/{row['id']}/schedule", {"month": first, "date": f"{first}-20"}).json["data"]
    assert moved["schedule"][first]["date"] == f"{first}-20"


def test_map_pins_follow_the_sites_state(app, client):
    row, month = visit_for(client)
    post(client, f"/contracts/{row['id']}", {})
    client.patch(f"/api/hub/contracts/{row['id']}", json={"latitude": 25.1, "longitude": 55.2}, headers=csrf(client))
    state = lambda: client.get("/api/hub/site-locations").json["data"][0]["state"]
    assert state() in ("Pending", "Approved")
    add_photos(client, row, month)
    complete(client, row, month)
    assert state() == "Pending Approval"
    post(client, f"/contracts/{row['id']}/reject", {"month": month, "reason": "Redo"})
    assert state() == "Rejected"


def test_staff_login_is_created_reset_and_closed_with_the_profile(app, client):
    person = post(client, "/staff", {"name": "Field Tester", "email": "field@example.test"}).json["data"]
    assert post(client, f"/staff/{person['id']}/login", {"password": "short"}).status_code == 400
    assert post(client, f"/staff/{person['id']}/login", {"password": "field-test-password", "role": "administrator"}).status_code == 400
    assert post(client, f"/staff/{person['id']}/login", {"password": "field-test-password", "role": "accountant"}).status_code == 400
    assert post(client, f"/staff/{person['id']}/login", {"password": "field-test-password"}).status_code == 201
    other = post(client, "/staff", {"name": "Twin", "email": "field@example.test"}).json["data"]
    assert post(client, f"/staff/{other['id']}/login", {"password": "field-test-password"}).status_code == 409
    def sign_in(password):
        return app.test_client().post("/api/auth/login", json={"email": "field@example.test", "password": password, "login_type": "staff"})
    assert sign_in("field-test-password").json["role"] == "employee"
    assert post(client, f"/staff/{person['id']}/login", {"password": "second-test-password"}).status_code == 200
    assert sign_in("second-test-password").status_code == 200
    patch(client, f"/staff/{person['id']}", {"status": "Resigned"})
    assert sign_in("second-test-password").status_code == 401
    patch(client, f"/staff/{person['id']}", {"status": "Active"})
    assert sign_in("second-test-password").status_code == 200
    client.delete(f"/api/hub/staff/{person['id']}", headers=csrf(client))
    assert sign_in("second-test-password").status_code == 401


def test_only_administrators_create_logins_and_staff_rename_and_picture_follow(app, client):
    coordinator = role_client(app, "service_coordinator", "Coordinator")
    person = post(client, "/staff", {"name": "Field Tester", "email": "field2@example.test"}).json["data"]
    assert post(coordinator, f"/staff/{person['id']}/login", {"password": "field-test-password"}).status_code == 403
    post(client, f"/staff/{person['id']}/login", {"password": "field-test-password"})
    staff = app.test_client()
    staff.csrf_token = staff.post("/api/auth/login", json={"email": "field2@example.test", "password": "field-test-password", "login_type": "staff"}).json["csrf_token"]
    assert staff.get("/api/hub/bootstrap").json["staff_id"] == person["id"]
    assert patch(staff, "/me/profile", {"full_name": "New Name"}).json["data"]["full_name"] == "New Name"
    assert client.get("/api/hub/staff").json["data"][0]["name"] == "New Name"
    assert upload(staff, "/me/photo", "me.png", PHOTO).status_code == 201
    entry = staff.get("/api/hub/file-index?collection=staff").json["data"][person["id"]]
    assert staff.get(f"/api/hub/files/{entry['photo']}?inline=1").status_code == 200
    assert upload(client, "/me/photo", "me.png", PHOTO).status_code == 404


def test_client_schedule_export_groups_sites_by_client(client):
    post(client, "/contracts", {"client": "Asteco", "site": "Madinat Badr"})
    post(client, "/contracts", {"client": "Asteco", "site": "Second Site"})
    from openpyxl import load_workbook
    sheet = load_workbook(BytesIO(client.get("/api/hub/export/client-schedule?format=xlsx").data)).active
    assert [cell.value for cell in sheet[1]][:2] == ["SL NO", "Client"]
    assert sheet["D2"].value == "Madinat Badr" and sheet["D3"].value == "Second Site"
    assert client.get("/api/hub/export/client-schedule?format=csv").data.startswith(b"\xef\xbb\xbfSL NO")
    assert client.get("/api/hub/export/client-schedule?format=pdf").data.startswith(b"%PDF")
    assert client.get("/api/hub/export/client-schedule?format=doc").status_code == 400


def test_old_sections_are_gone(client):
    for path in ("/api/projects", "/api/engineers", "/api/water-testing", "/api/cleaning-records", "/api/certificate-templates", "/api/approvals", "/api/uploads", "/api/reports/export"):
        assert client.get(path).status_code == 404, path


def test_projects_and_engineers_move_to_the_new_sections_before_their_tables_go():
    import json
    import sqlite3
    from pathlib import Path
    folder = Path(__file__).parent.parent / "d1-migrations"
    connection = sqlite3.connect(":memory:")
    for migration in sorted(folder.glob("*.sql")):
        if migration.name.startswith("0004"):
            connection.executescript("""
                INSERT INTO organizations (id, name, created_at, updated_at) VALUES ('org', 'Wolgan', 'now', 'now');
                INSERT INTO clients (id, organization_id, name, is_active, created_at, updated_at) VALUES ('c1', 'org', 'Daikin', 1, 'now', 'now');
                INSERT INTO projects (id, organization_id, client_id, name, so_no, engineer_visit, start_date, end_date, latitude, longitude, service, status, created_at, updated_at)
                    VALUES ('p1', 'org', 'c1', 'True Bell', 1087, 'Quarterly', '2026-01-01', '2026-12-31', 25.1, 55.2, 'Water testing', 'Planned', 'now', 'now');
                INSERT INTO projects (id, organization_id, client_id, name, status, deleted_at, created_at, updated_at)
                    VALUES ('p2', 'org', 'c1', 'Deleted', 'Planned', 'now', 'now', 'now');
                INSERT INTO projects (id, organization_id, client_id, name, engineer_name, engineer_visit, engineer_visit_custom_date, status, created_at, updated_at)
                    VALUES ('p3', 'org', 'c1', 'JTM', 'Shahul', 'Custom', '2026-10-29', 'Planned', 'now', 'now');
                INSERT INTO engineers (id, organization_id, full_name, engineer_code, is_active, created_at, updated_at) VALUES ('e1', 'org', 'Ali', 'E1', 1, 'now', 'now');
            """)
        connection.executescript(migration.read_text(encoding="utf-8"))
    rows = connection.execute("SELECT collection, data FROM hub_records ORDER BY collection, id").fetchall()
    assert [row[0] for row in rows] == ["contracts", "contracts", "staff"]
    moved = json.loads(rows[0][1])
    assert moved["site"] == "True Bell" and moved["client"] == "Daikin" and moved["so"] == "1087"
    assert moved["frequency"] == "Quarterly" and moved["latitude"] == 25.1 and moved["schedule"] == {}
    custom = json.loads(rows[1][1])
    assert custom["schedule"] == {"2026-10": {"status": "due", "date": "2026-10-29", "engineer": "Shahul"}}
    assert custom["next_due"] == "2026-10-29" and custom["frequency"] == ""
    assert json.loads(rows[2][1])["name"] == "Ali"
    tables = {row[0] for row in connection.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
    assert not tables & {"projects", "engineers", "certificate_templates", "water_testing_records", "cleaning_records"}


def test_backup_is_for_administrators(app, client):
    contract(client)
    assert client.get("/api/hub/backup").status_code == 200
    assert role_client(app, "accountant", "Accounts").get("/api/hub/backup").status_code == 403


def test_overview_counts(client):
    contract(client)
    contract(client, site="Old", expiry="2026-01-01", start="2025-01-01")
    closed = contract(client, site="Closed")
    post(client, f"/contracts/{closed['id']}/close")
    stats = client.get("/api/hub/overview").json
    assert stats["contracts"]["active"] == 2 and stats["contracts"]["expired"] == 1
