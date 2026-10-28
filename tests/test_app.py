import os
import shutil
import subprocess
import time
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from werkzeug.security import generate_password_hash

from backend import create_app
from backend.database import db
from backend.models import clients, engineers, memberships, organizations, projects, users
from backend.routes.api import safe_spreadsheet_value
from backend.utils.validators import validate_payload


ADMIN_KEY = "private-admin-key-1"


USE_D1 = os.environ.get("WATERCARE_TEST_D1") == "1"


class TestConfig:
    TESTING = True
    SECRET_KEY = "test-secret-key"
    # WATERCARE_TEST_D1=1 runs every test through the Cloudflare D1 driver, backed by an in-memory D1 stand-in.
    SQLALCHEMY_DATABASE_URI = "sqlite+d1://" if USE_D1 else "sqlite://"
    SQLALCHEMY_TRACK_MODIFICATIONS = False
    STORAGE_PATH = "test-uploads"


@pytest.fixture
def app():
    if USE_D1:
        from backend import d1
        d1.use_runner(d1.SQLiteRunner())
    application = create_app(TestConfig)
    with application.app_context():
        db.create_all()
        now = datetime.now(timezone.utc).isoformat()
        db.session.execute(organizations.insert().values(
            id="org-1", name="Test organization", created_at=now, updated_at=now,
        ))
        db.session.execute(users.insert().values(
            id="user-1", email="admin@example.test", full_name="Admin",
            password_hash=generate_password_hash("secure-test-password"),
            access_key_hash=generate_password_hash(ADMIN_KEY),
            is_active=True, created_at=now, updated_at=now,
        ))
        db.session.execute(memberships.insert().values(
            id=str(uuid4()), organization_id="org-1", user_id="user-1",
            role="administrator", status="active", created_at=now,
        ))
        db.session.commit()
        yield application
        db.drop_all()


@pytest.fixture
def client(app):
    test_client = app.test_client()
    response = test_client.post("/api/auth/login", json={
        "email": "admin@example.test", "password": "secure-test-password",
        "login_type": "admin", "access_key": ADMIN_KEY,
    })
    assert response.status_code == 200
    test_client.csrf_token = response.json["csrf_token"]
    return test_client


def csrf(client):
    return {"X-CSRF-Token": client.csrf_token}


def test_health_checks_database(app):
    assert app.test_client().get("/api/health").json == {"status": "ok"}


def test_production_requires_environment_configuration(monkeypatch):
    class MissingConfig:
        TESTING = False
        SECRET_KEY = ""
        SQLALCHEMY_DATABASE_URI = "sqlite://"
        STORAGE_PATH = "test-uploads"

    monkeypatch.setenv("FLASK_ENV", "production")
    monkeypatch.delenv("SECRET_KEY", raising=False)
    monkeypatch.delenv("DATABASE_URL", raising=False)
    with pytest.raises(RuntimeError, match="SECRET_KEY, DATABASE_URL"):
        create_app(MissingConfig)
    monkeypatch.delenv("FLASK_ENV")


def test_login_logout_and_password_session(client):
    assert client.get("/api/auth/me").json["role"] == "administrator"
    response = client.post("/api/auth/logout", headers=csrf(client))
    assert response.status_code == 200
    assert client.get("/api/auth/me").status_code == 401


def test_wrong_password_is_rejected(app):
    response = app.test_client().post("/api/auth/login", json={
        "email": "admin@example.test", "password": "incorrect",
    })
    assert response.status_code == 401


def test_authenticated_user_can_create_and_read_client(client):
    response = client.post("/api/clients", headers=csrf(client), json={"name": "Acme"})
    assert response.status_code == 201
    item_id = response.json["data"]["id"]
    assert client.get(f"/api/clients/{item_id}").json["data"]["name"] == "Acme"


def test_administrator_can_create_an_employee_account(client):
    response = client.post("/api/auth/users", headers=csrf(client), json={
        "email": "new@example.test", "full_name": "New Employee",
        "password": "a-long-test-password", "role": "employee",
    })
    assert response.status_code == 201
    employee = client.application.test_client()
    login_response = employee.post("/api/auth/login", json={
        "email": "new@example.test", "password": "a-long-test-password",
    })
    assert login_response.status_code == 200
    assert login_response.json["role"] == "employee"
    engineers_response = client.get("/api/engineers")
    assert any(engineer["user_id"] == response.json["data"]["id"] for engineer in engineers_response.json["data"])


def test_csrf_and_authentication_are_required(app):
    anonymous = app.test_client()
    assert anonymous.get("/api/clients").status_code == 401
    assert anonymous.post("/api/clients", json={"name": "Acme"}).status_code == 401
    client = app.test_client()
    client.post("/api/auth/login", json={"email": "admin@example.test", "password": "secure-test-password", "login_type": "admin", "access_key": ADMIN_KEY})
    assert client.post("/api/clients", json={"name": "Acme"}).status_code == 403


def test_employee_cannot_create_client(app):
    with app.app_context():
        now = datetime.now(timezone.utc).isoformat()
        db.session.execute(users.insert().values(
            id="employee-1", email="employee@example.test", full_name="Employee",
            password_hash=generate_password_hash("secure-test-password"),
            is_active=True, created_at=now, updated_at=now,
        ))
        db.session.execute(memberships.insert().values(
            id=str(uuid4()), organization_id="org-1", user_id="employee-1",
            role="employee", status="active", created_at=now,
        ))
        db.session.commit()
    client = app.test_client()
    response = client.post("/api/auth/login", json={
        "email": "employee@example.test", "password": "secure-test-password",
    })
    assert response.status_code == 200
    assert client.post("/api/clients", headers={"X-CSRF-Token": response.json["csrf_token"]}, json={"name": "Blocked"}).status_code == 403


def test_user_cannot_access_records_from_another_organization(client):
    with client.session_transaction() as browser_session:
        browser_session["organization_id"] = "another-org"
    assert client.get("/api/clients").status_code == 401


def test_private_document_upload_download_and_delete(client):
    created = client.post("/api/clients", headers=csrf(client), json={"name": "Document client"})
    client_id = created.json["data"]["id"]
    upload = client.post(
        "/api/uploads",
        headers=csrf(client),
        data={
            "related_type": "client",
            "related_id": client_id,
            "file": (BytesIO(b"%PDF-1.4\nWaterCare test\n%%EOF"), "result.pdf", "application/pdf"),
        },
        content_type="multipart/form-data",
    )
    assert upload.status_code == 201
    file_id = upload.json["data"]["id"]
    download = client.get(f"/api/uploads/{file_id}")
    assert download.status_code == 200
    assert download.data.startswith(b"%PDF-")
    download.close()
    assert client.delete(f"/api/uploads/{file_id}", headers=csrf(client)).status_code == 200
    assert client.get(f"/api/uploads/{file_id}").status_code == 404


def test_project_entry_completion_durations_and_multiple_images(client):
    created = client.post("/api/projects", headers=csrf(client), json={
        "client_name": "WaterCare Test Client",
        "name": "Treatment System",
        "service": "Water testing",
        "so_no": 1234,
        "start_date": "2026-12-30",
        "visit_time": "08:30",
        "end_date": "2026-12-31",
        "system_details": "Main plant",
        "engineer_name": "Assigned Engineer",
        "engineer_visit": "Quarterly",
        "microbiology": "No",
        "microbiology_frequency": "No",
        "legionella": "Yes",
        "legionella_frequency": "Alternative",
    })
    assert created.status_code == 201
    project_id = created.json["data"]["id"]
    assert created.json["data"]["client_name"] == "WaterCare Test Client"
    assert created.json["data"]["service"] == "Water testing"
    calendar = client.get("/api/dashboard/tasks?from=2026-12-30&to=2026-12-31")
    assert any(
        task.get("project", {}).get("service") == "Water testing"
        for task in calendar.json["data"]
    )
    project_days = [
        task["scheduled_date"] for task in calendar.json["data"]
        if task.get("project", {}).get("id") == project_id
    ]
    assert project_days == ["2026-12-30", "2026-12-31"]
    assert next(task for task in calendar.json["data"] if task["scheduled_date"] == "2026-12-30")["scheduled_time"] == "08:30"

    upload_completion_images(client, project_id)
    completed = client.patch(f"/api/projects/{project_id}", headers=csrf(client), json={
        "completion_date": "2026-12-31",
        "completion_engineer_name": "Assigned Engineer",
        "invoice_number": "INV-1234",
    })
    assert completed.status_code == 200, completed.json
    assert completed.json["data"]["approval_status"] == "Pending Approval"
    assert completed.json["data"]["status"] != "Completed"
    assert completed.json["data"]["completion_date"] == "2026-12-31"
    assert completed.json["data"]["completion_engineer_name"] == "Assigned Engineer"

    for image_type, filename in (("before", "before.png"), ("after", "after.png")):
        uploaded = client.post(
            "/api/uploads",
            headers=csrf(client),
            data={
                "related_type": "project",
                "related_id": project_id,
                "image_type": image_type,
                "file": (BytesIO(b"\x89PNG\r\n\x1a\nimage"), filename, "image/png"),
            },
            content_type="multipart/form-data",
        )
        assert uploaded.status_code == 201
        assert uploaded.json["data"]["image_type"] == image_type

    images = client.get(f"/api/uploads?related_type=project&related_id={project_id}")
    assert {item["image_type"] for item in images.json["data"]} == {"before", "after"}


def test_named_engineer_can_complete_assigned_project(app, client):
    project = client.post("/api/projects", headers=csrf(client), json={
        "client_name": "Engineer Client", "name": "Engineer Project", "service": "Tank cleaning",
        "contract_period": "2026-12-30", "engineer_name": "Field Engineer",
        "engineer_visit": "Quarterly", "microbiology": "No",
        "microbiology_frequency": "No", "legionella": "No",
        "legionella_frequency": "No",
    }).json["data"]
    with app.app_context():
        now = datetime.now(timezone.utc).isoformat()
        db.session.execute(users.insert().values(
            id="employee-2", email="field@example.test", full_name="Field Engineer",
            password_hash=generate_password_hash("secure-test-password"),
            is_active=True, created_at=now, updated_at=now,
        ))
        db.session.execute(memberships.insert().values(
            id=str(uuid4()), organization_id="org-1", user_id="employee-2",
            role="employee", status="active", created_at=now,
        ))
        db.session.execute(engineers.insert().values(
            id="engineer-2", organization_id="org-1", full_name="Field Engineer",
            user_id="employee-2", is_active=True, created_at=now, updated_at=now,
        ))
        db.session.commit()

    engineer = app.test_client()
    login = engineer.post("/api/auth/login", json={
        "email": "field@example.test", "password": "secure-test-password",
    })
    engineer.csrf_token = login.json["csrf_token"]
    visible = engineer.get("/api/projects")
    assert visible.status_code == 200
    assert any(item["id"] == project["id"] for item in visible.json["data"])
    upload_completion_images(engineer, project["id"])
    completed = engineer.patch(f"/api/projects/{project['id']}", headers=csrf(engineer), json={
        "completion_date": "2026-12-30",
        "completion_engineer_name": "Field Engineer",
        "invoice_number": "INV-ENGINEER",
    })
    assert completed.status_code == 200
    for image_type, filename in (("before", "employee-before.png"), ("after", "employee-after.png")):
        uploaded = engineer.post(
            "/api/uploads",
            headers=csrf(engineer),
            data={
                "related_type": "project",
                "related_id": project["id"],
                "image_type": image_type,
                "file": (BytesIO(b"\x89PNG\r\n\x1a\nimage"), filename, "image/png"),
            },
            content_type="multipart/form-data",
        )
        assert uploaded.status_code == 201


def test_permanent_project_delete_removes_related_documents(client):
    project = client.post("/api/projects", headers=csrf(client), json={
        "client_name": "Delete Project Client",
        "name": "Project to delete",
        "service": "Tank cleaning",
        "so_no": 9001,
        "start_date": "2026-10-24",
        "end_date": "2026-10-25",
        "location": "Test site",
        "system_details": "Test system",
        "engineer_name": "Test Engineer",
        "engineer_visit": "Quarterly",
        "microbiology": "No",
        "microbiology_frequency": "No",
        "legionella": "No",
        "legionella_frequency": "No",
    }).json["data"]
    upload = client.post(
        "/api/uploads",
        headers=csrf(client),
        data={
            "related_type": "project",
            "related_id": project["id"],
            "image_type": "before",
            "file": (BytesIO(b"\x89PNG\r\n\x1a\nimage"), "before.png", "image/png"),
        },
        content_type="multipart/form-data",
    )
    assert upload.status_code == 201
    file_id = upload.json["data"]["id"]
    path = Path(client.application.config["STORAGE_PATH"]) / upload.json["data"]["storage_path"]
    assert path.is_file()

    deleted = client.delete(f"/api/projects/{project['id']}/permanent", headers=csrf(client))
    assert deleted.status_code == 200
    assert client.get(f"/api/projects/{project['id']}").status_code == 404
    assert client.get(f"/api/uploads/{file_id}").status_code == 404
    assert client.get(f"/api/uploads?related_type=project&related_id={project['id']}").json["data"] == []
    assert not path.exists()


def test_calendar_rejects_malformed_dates(client):
    response = client.get("/api/dashboard/tasks?from=20261023&to=2026-10-24")
    assert response.status_code == 400


def test_impossible_project_dates_are_rejected_and_do_not_break_the_calendar(app, client):
    project = {"client_name": "Date Client", "name": "Date Project", "service": "Water testing",
               "start_date": "2026-11-31", "end_date": "2026-12-02"}
    rejected = client.post("/api/projects", headers=csrf(client), json=project)
    assert rejected.status_code == 400
    assert "real calendar date" in rejected.json["error"]["message"]

    project["start_date"] = "2026-11-28"
    valid = client.post("/api/projects", headers=csrf(client), json=project).json["data"]
    legacy = client.post("/api/projects", headers=csrf(client), json={**project, "name": "Legacy Project"}).json["data"]
    with app.app_context():
        db.session.execute(projects.update().where(projects.c.id == legacy["id"]).values(start_date="2026-11-31"))
        db.session.commit()
    calendar = client.get("/api/dashboard/tasks?from=2026-11-01&to=2026-11-30")
    assert calendar.status_code == 200
    shown = {task["id"] for task in calendar.json["data"]}
    assert valid["id"] in shown and legacy["id"] not in shown


def run_browser_test(filename):
    node = shutil.which("node")
    if not node:
        pytest.skip("Node.js is required for the browser form tests.")
    result = subprocess.run(
        [node, "--test", str(Path(__file__).parent / "js" / filename)],
        stdin=subprocess.DEVNULL, capture_output=True, text=True, timeout=60,
    )
    assert result.returncode == 0, result.stdout + result.stderr


def test_water_testing_form_round_trips_json_fields():
    run_browser_test("json_fields.test.js")


def test_map_location_search_suggestions_and_typed_coordinates():
    run_browser_test("place_search.test.js")


def test_wrong_http_method_returns_json_405(client):
    response = client.put("/api/auth/login", json={})
    assert response.status_code == 405
    assert response.json["error"]["code"] == "method_not_allowed"
    assert "POST" in response.headers["Allow"]


def test_rejected_project_does_not_create_its_client(app, client):
    def client_count():
        with app.app_context():
            return db.session.execute(db.select(db.func.count()).select_from(clients)).scalar()

    before = client_count()
    rejected = client.post("/api/projects", headers=csrf(client), json={
        "client_name": "Orphan Client", "name": "Rejected Project", "service": "Water testing",
        "assigned_engineer_ids": [str(uuid4())],
    })
    assert rejected.status_code == 400
    missing = client.patch(f"/api/projects/{uuid4()}", headers=csrf(client), json={"client_name": "Orphan Client"})
    assert missing.status_code == 404
    assert client_count() == before


def test_settings_are_validated_and_record_each_update(client):
    for invalid in (
        {"user": {"appearance": "x" * 500}},
        {"user": {"time_format": "25h"}},
        {"user": {"time_zone": {"zone": "UTC"}}},
        {"user": {"notification_preferences": "email"}},
        {"organization": {"name": 7}},
    ):
        response = client.patch("/api/settings", headers=csrf(client), json=invalid)
        assert response.status_code == 400, invalid

    assert client.patch("/api/settings", headers=csrf(client), json={"user": {"appearance": "dark"}}).status_code == 200
    first = client.get("/api/settings").json["user"]["updated_at"]
    time.sleep(0.01)
    assert client.patch("/api/settings", headers=csrf(client), json={"user": {"appearance": "light"}}).status_code == 200
    second = client.get("/api/settings").json["user"]
    assert second["appearance"] == "light"
    assert second["updated_at"] > first


def test_settings_are_saved_with_organization_and_user_scope(client):
    response = client.patch("/api/settings", headers=csrf(client), json={
        "organization": {"name": "Updated organization"},
        "user": {"appearance": "dark", "time_zone": "UTC"},
    })
    assert response.status_code == 200
    saved = client.get("/api/settings").json
    assert saved["organization"]["name"] == "Updated organization"
    assert saved["user"]["appearance"] == "dark"


def test_login_shell_preserves_watercare_brand_and_has_no_provider_globals(app):
    response = app.test_client().get("/")
    assert response.status_code == 200
    assert b"WaterCare" in response.data
    assert b"window.WATERCARE_" not in response.data


def test_payload_validation_and_spreadsheet_safety():
    assert "name" in validate_payload("clients", {})
    assert validate_payload("clients", {"name": "WaterCare"}) is None
    assert "Unsupported field" in validate_payload("clients", {"name": "Client", "organization_id": "attacker"})
    assert safe_spreadsheet_value("=1+1") == "'=1+1"
    assert safe_spreadsheet_value("-2") == "'-2"
    assert safe_spreadsheet_value(2) == 2
    base_project = {
        "client_name": "Acme", "name": "System", "service": "Water testing", "start_date": "2026-10-25",
        "end_date": "2026-10-24", "microbiology": "Yes",
        "microbiology_frequency": "Quarterly", "legionella": "No",
        "legionella_frequency": "No",
    }
    assert "on or after" in validate_payload("projects", base_project)
    base_project["end_date"] = "2026-10-26"
    assert validate_payload("projects", base_project) is None
    base_project["microbiology_frequency"] = "No"
    assert "required when" in validate_payload("projects", base_project)
    base_project.update({"microbiology_frequency": "Alternative", "visit_time": "09:30"})
    assert validate_payload("projects", base_project) is None


# --- Engineers, schedule frequencies, approvals, roles, invoice numbers and map locations ---

ROLE_PASSWORD = "role-test-password-1"


ROLE_KEY = "role-private-key-1"


def sign_in(app, email, password=ROLE_PASSWORD, key=None):
    test_client = app.test_client()
    body = {"email": email, "password": password}
    if key:
        body.update(login_type="admin", access_key=key)
    response = test_client.post("/api/auth/login", json=body)
    assert response.status_code == 200, response.json
    test_client.csrf_token = response.json["csrf_token"]
    return test_client


def add_account(app, admin, role, name):
    email = f"{role.replace('_', '-')}-{uuid4().hex[:6]}@example.test"
    key = ROLE_KEY if role in {"service_coordinator", "accountant"} else None
    created = admin.post("/api/auth/users", headers=csrf(admin), json={
        "email": email, "full_name": name, "password": ROLE_PASSWORD, "role": role, **({"access_key": key} if key else {}),
    })
    assert created.status_code == 201, created.json
    return sign_in(app, email, key=key)


def new_project(test_client, **values):
    payload = {
        "client_name": "Gulf Towers", "name": "Cooling Tower", "service": "Water testing", "so_no": 501,
        "start_date": "2026-11-10", "end_date": "2026-11-11", "location": "Tower A",
        "system_details": "Two cooling towers", "engineer_name": "Unassigned",
        "engineer_visit": "Quarterly", "microbiology": "No", "microbiology_frequency": "No",
        "legionella": "No", "legionella_frequency": "No", **values,
    }
    response = test_client.post("/api/projects", headers=csrf(test_client), json=payload)
    assert response.status_code == 201, response.json
    return response.json["data"]


def upload_completion_images(test_client, project_id):
    for image_type in ("before", "after"):
        uploaded = test_client.post(
            "/api/uploads",
            headers=csrf(test_client),
            data={
                "related_type": "project",
                "related_id": project_id,
                "image_type": image_type,
                "file": (BytesIO(b"\x89PNG\r\n\x1a\nimage"), f"{image_type}.png", "image/png"),
            },
            content_type="multipart/form-data",
        )
        assert uploaded.status_code == 201, uploaded.json


def employee_with_engineer(app, admin, name, code):
    employee = add_account(app, admin, "employee", name)
    profile = next(item for item in admin.get("/api/engineers").json["data"] if item["full_name"] == name)
    assert admin.patch(f"/api/engineers/{profile['id']}", headers=csrf(admin), json={"engineer_code": code}).status_code == 200
    return employee, profile


def test_service_coordinator_and_accountant_edit_while_other_roles_are_view_only(app, client):
    assert client.post("/api/auth/users", headers=csrf(client), json={
        "email": "old-manager@example.test", "full_name": "Old Manager", "password": ROLE_PASSWORD, "role": "manager",
    }).status_code == 400
    coordinator = add_account(app, client, "service_coordinator", "Sara Coordinator")
    accountant = add_account(app, client, "accountant", "Omar Accountant")
    employee = add_account(app, client, "employee", "Eli Employee")
    viewer = add_account(app, client, "viewer", "Vera Viewer")
    assert coordinator.get("/api/auth/me").json["role"] == "service_coordinator"

    for editor in (coordinator, accountant):
        assert editor.post("/api/clients", headers=csrf(editor), json={"name": f"Client {uuid4().hex[:4]}"}).status_code == 201
        assert editor.post("/api/engineers", headers=csrf(editor), json={
            "full_name": f"Engineer {uuid4().hex[:4]}", "engineer_code": f"E-{uuid4().hex[:6]}",
        }).status_code == 201
        project = new_project(editor)
        assert editor.patch(f"/api/projects/{project['id']}", headers=csrf(editor), json={"location": "Tower B"}).status_code == 200
        assert editor.post("/api/auth/users", headers=csrf(editor), json={
            "email": f"x{uuid4().hex[:4]}@example.test", "full_name": "X", "password": ROLE_PASSWORD, "role": "employee",
        }).status_code == 403

    project = new_project(coordinator)
    client_id = coordinator.get("/api/clients").json["data"][0]["id"]
    for view_only in (employee, viewer):
        headers = csrf(view_only)
        assert view_only.post("/api/clients", headers=headers, json={"name": "Blocked"}).status_code == 403
        assert view_only.post("/api/projects", headers=headers, json={"client_name": "Blocked", "name": "P", "service": "Water testing"}).status_code == 403
        assert view_only.post("/api/engineers", headers=headers, json={"full_name": "Blocked", "engineer_code": "B-1"}).status_code == 403
        assert view_only.patch(f"/api/projects/{project['id']}", headers=headers, json={"name": "Renamed"}).status_code == 403
        assert view_only.patch(f"/api/clients/{client_id}", headers=headers, json={"name": "Renamed"}).status_code == 403
        assert view_only.delete(f"/api/projects/{project['id']}", headers=headers).status_code == 403
        assert view_only.delete(f"/api/projects/{project['id']}/permanent", headers=headers).status_code == 403
        assert view_only.post(f"/api/projects/{project['id']}/approve", headers=headers).status_code == 403
        assert view_only.post(f"/api/projects/{project['id']}/reject", headers=headers, json={"reason": "No"}).status_code == 403
    assert viewer.get("/api/projects").status_code == 200
    assert viewer.patch(f"/api/projects/{project['id']}", headers=csrf(viewer), json={
        "completion_status": "Completed Within Contract Period",
    }).status_code == 403
    assert coordinator.get(f"/api/projects/{project['id']}").json["data"]["name"] == "Cooling Tower"


def test_frequency_options_custom_dates_and_existing_values(app, client):
    project = new_project(
        client, engineer_visit="Custom", engineer_visit_custom_date="2026-12-05",
        microbiology="Yes", microbiology_frequency="Alternative",
        legionella="Yes", legionella_frequency="Custom", legionella_custom_date="2027-01-15",
    )
    saved = client.get(f"/api/projects/{project['id']}").json["data"]
    assert (saved["engineer_visit"], saved["engineer_visit_custom_date"]) == ("Custom", "2026-12-05")
    assert (saved["microbiology_frequency"], saved["microbiology_custom_date"]) == ("Alternative", None)
    assert (saved["legionella_frequency"], saved["legionella_custom_date"]) == ("Custom", "2027-01-15")

    headers = csrf(client)
    for invalid, message in (
        ({"engineer_visit": "Custom"}, "Choose a date"),
        ({"engineer_visit": "Quarterly", "engineer_visit_custom_date": "2026-12-05"}, "only be set when"),
        ({"engineer_visit": "Monthly"}, "Unsupported engineer visit"),
        ({"microbiology": "No", "microbiology_frequency": "Quarterly"}, "must be No"),
        ({"microbiology": "Yes", "microbiology_frequency": "Weekly"}, "required when"),
        ({"legionella_custom_date": "2027-02-30"}, "real calendar date"),
    ):
        response = client.post("/api/projects", headers=headers, json={
            "client_name": "Freq Client", "name": "Freq", "service": "Water testing", **invalid,
        })
        assert response.status_code == 400 and message in response.json["error"]["message"], (invalid, response.json)

    switched = client.patch(f"/api/projects/{project['id']}", headers=headers, json={"engineer_visit": "Quarterly"})
    assert switched.status_code == 200
    assert switched.json["data"]["engineer_visit_custom_date"] is None
    assert client.patch(f"/api/projects/{project['id']}", headers=headers, json={"legionella_custom_date": None}).status_code == 400

    with app.app_context():
        db.session.execute(projects.update().where(projects.c.id == project["id"]).values(
            engineer_visit="Monthly", microbiology_frequency="Weekly",
        ))
        db.session.commit()
    kept = client.patch(f"/api/projects/{project['id']}", headers=headers, json={
        "name": "Cooling Tower North", "engineer_visit": "Monthly", "microbiology": "Yes", "microbiology_frequency": "Weekly",
    })
    assert kept.status_code == 200, kept.json
    assert (kept.json["data"]["engineer_visit"], kept.json["data"]["microbiology_frequency"]) == ("Monthly", "Weekly")
    assert client.patch(f"/api/projects/{project['id']}", headers=headers, json={"engineer_visit": "Weekly"}).status_code == 400


def test_engineers_have_unique_ids_and_can_be_selected_for_projects(app, client):
    headers = csrf(client)
    assert client.post("/api/engineers", headers=headers, json={"full_name": "No Code"}).status_code == 400
    first = client.post("/api/engineers", headers=headers, json={"full_name": "Ali Hassan", "engineer_code": " ENG-001 "})
    assert first.status_code == 201 and first.json["data"]["engineer_code"] == "ENG-001"
    assert client.post("/api/engineers", headers=headers, json={"full_name": "Copy", "engineer_code": "eng-001"}).status_code == 409
    second = client.post("/api/engineers", headers=headers, json={"full_name": "Mona Saleh", "engineer_code": "ENG-002"}).json["data"]
    assert client.patch(f"/api/engineers/{second['id']}", headers=headers, json={"engineer_code": "ENG-001"}).status_code == 409
    assert client.patch(f"/api/engineers/{second['id']}", headers=headers, json={"engineer_code": "ENG-002", "phone": "555"}).status_code == 200
    listed = {item["engineer_code"]: item["full_name"] for item in client.get("/api/engineers").json["data"]}
    assert listed["ENG-001"] == "Ali Hassan" and listed["ENG-002"] == "Mona Saleh"

    project = new_project(client, engineer_name="Ali Hassan", assigned_engineer_ids=[first.json["data"]["id"]])
    row = next(item for item in client.get("/api/projects").json["data"] if item["id"] == project["id"])
    assert row["assigned_engineer_ids"] == [first.json["data"]["id"]]

    employee, profile = employee_with_engineer(app, client, "Field Engineer Two", "ENG-003")
    assert employee.get("/api/projects").json["data"] == []
    assert client.patch(f"/api/projects/{project['id']}", headers=headers, json={
        "engineer_name": "Field Engineer Two", "assigned_engineer_ids": [profile["id"]],
    }).status_code == 200
    assert [item["id"] for item in employee.get("/api/projects").json["data"]] == [project["id"]]


def test_date_filter_returns_only_work_recorded_for_the_selected_date(client):
    selected = new_project(client, name="Selected Day Project", start_date="2026-11-10", end_date="2026-11-11")
    other = new_project(client, name="Other Day Project", start_date="2026-11-12", end_date="2026-11-12")
    client_id = selected["client_id"]
    cleaning = client.post("/api/cleaning-records", headers=csrf(client), json={
        "client_id": client_id, "cleaning_date": "2026-11-10", "tank_identification": "Tank 4",
    }).json["data"]
    testing = client.post("/api/water-testing", headers=csrf(client), json={
        "client_id": client_id, "collection_date": "2026-11-11", "sample_identification": "S-9",
    }).json["data"]

    items = client.get("/api/dashboard/tasks?from=2026-11-10&to=2026-11-10").json["data"]
    ids = {item["id"] for item in items}
    assert {selected["id"], cleaning["id"]} <= ids
    assert other["id"] not in ids and testing["id"] not in ids
    assert all(item["scheduled_date"] == "2026-11-10" for item in items)
    project_item = next(item for item in items if item["id"] == selected["id"])
    assert project_item["project"]["client_name"] == "Gulf Towers"
    assert project_item["project"]["name"] == "Selected Day Project"


def test_invoice_number_is_saved_with_the_completion(client):
    project = new_project(client)
    headers = csrf(client)
    upload_completion_images(client, project["id"])
    too_long = client.patch(f"/api/projects/{project['id']}", headers=headers, json={
        "completion_date": "2026-11-10",
        "completion_engineer_name": "Engineer", "invoice_number": "X" * 101,
    })
    assert too_long.status_code == 400
    submitted = client.patch(f"/api/projects/{project['id']}", headers=headers, json={
        "completion_date": "2026-11-10",
        "completion_engineer_name": "Engineer", "invoice_number": "  INV-2026-0042 ",
    })
    assert submitted.status_code == 200
    assert client.get(f"/api/projects/{project['id']}").json["data"]["invoice_number"] == "INV-2026-0042"
    corrected = client.patch(f"/api/projects/{project['id']}", headers=headers, json={"invoice_number": "INV-2026-0043"})
    assert corrected.status_code == 200
    assert corrected.json["data"]["invoice_number"] == "INV-2026-0043"
    assert corrected.json["data"]["approval_status"] == "Pending Approval"


def test_completion_requires_photos_and_derives_timing_from_actual_date(client):
    project = new_project(client)
    url = f"/api/projects/{project['id']}"
    payload = {
        "completion_date": "2026-11-09",
        "completion_engineer_name": "Engineer",
        "invoice_number": "INV-1",
        "completion_status": "Completed After the Scheduled Date",
    }
    missing_photos = client.patch(url, headers=csrf(client), json=payload)
    assert missing_photos.status_code == 400
    assert "both before-work and after-work photos" in missing_photos.json["error"]["message"]

    upload_completion_images(client, project["id"])
    submitted = client.patch(url, headers=csrf(client), json=payload)
    assert submitted.status_code == 200
    assert submitted.json["data"]["approval_status"] == "Pending Approval"
    assert submitted.json["data"]["completion_status"] == "Completed Before the Scheduled Date"

    for index, (actual_date, expected_status) in enumerate((
        ("2026-11-10", "Completed on the Scheduled Date"),
        ("2026-11-11", "Completed After the Scheduled Date"),
    ), start=1):
        same_day_project = new_project(client, name=f"Completion timing {index}", so_no=600 + index)
        upload_completion_images(client, same_day_project["id"])
        result = client.patch(f"/api/projects/{same_day_project['id']}", headers=csrf(client), json={
            "completion_date": actual_date,
            "completion_engineer_name": "Engineer",
            "invoice_number": f"INV-{index}",
        })
        assert result.status_code == 200
        assert result.json["data"]["completion_status"] == expected_status


def test_completion_submission_approval_rejection_and_history(app, client):
    coordinator = add_account(app, client, "service_coordinator", "Sara Coordinator")
    accountant = add_account(app, client, "accountant", "Omar Accountant")
    employee, profile = employee_with_engineer(app, client, "Field Engineer", "ENG-010")
    project = new_project(coordinator, engineer_name="Field Engineer", assigned_engineer_ids=[profile["id"]], so_no=777)
    assert project["approval_status"] == "Pending"
    url = f"/api/projects/{project['id']}"

    upload_completion_images(employee, project["id"])
    submitted = employee.patch(url, headers=csrf(employee), json={
        "completion_date": "2026-11-11",
        "completion_engineer_name": "Field Engineer", "invoice_number": "INV-1",
    })
    assert submitted.status_code == 200, submitted.json
    assert submitted.json["data"]["approval_status"] == "Pending Approval"
    assert submitted.json["data"]["completion_date"] == "2026-11-11"
    assert submitted.json["data"]["status"] != "Completed"
    assert employee.patch(url, headers=csrf(employee), json={
        "completion_date": "2026-11-09",
        "completion_engineer_name": "Field Engineer", "invoice_number": "INV-1",
    }).status_code == 409
    assert employee.post(f"{url}/approve", headers=csrf(employee)).status_code == 403
    assert [row["id"] for row in accountant.get("/api/approvals").json["data"]] == [project["id"]]

    assert accountant.post(f"{url}/reject", headers=csrf(accountant), json={"reason": "  "}).status_code == 400
    rejected = accountant.post(f"{url}/reject", headers=csrf(accountant), json={"reason": "After photos are missing."})
    assert rejected.status_code == 200
    assert rejected.json["data"]["approval_status"] == "Rejected"
    assert employee.get(url).json["data"]["rejection_reason"] == "After photos are missing."
    assert accountant.post(f"{url}/approve", headers=csrf(accountant)).status_code == 409

    resubmitted = employee.patch(url, headers=csrf(employee), json={
        "completion_date": "2026-11-11",
        "completion_engineer_name": "Field Engineer", "invoice_number": "INV-2",
    })
    assert resubmitted.status_code == 200
    assert resubmitted.json["data"]["approval_status"] == "Pending Approval"
    assert resubmitted.json["data"]["rejection_reason"] is None
    approved = coordinator.post(f"{url}/approve", headers=csrf(coordinator))
    assert approved.status_code == 200
    assert (approved.json["data"]["approval_status"], approved.json["data"]["status"]) == ("Approved", "Completed")
    assert coordinator.post(f"{url}/approve", headers=csrf(coordinator)).status_code == 409
    assert employee.patch(url, headers=csrf(employee), json={
        "completion_date": "2026-11-10",
        "completion_engineer_name": "Field Engineer", "invoice_number": "INV-2",
    }).status_code == 409

    history = coordinator.get(f"/api/approval-history?project_id={project['id']}").json["data"]
    assert [(row["action"], row["previous_status"], row["new_status"], row["actor_name"]) for row in history] == [
        ("Approved", "Pending Approval", "Approved", "Sara Coordinator"),
        ("Resubmitted", "Rejected", "Pending Approval", "Field Engineer"),
        ("Rejected", "Pending Approval", "Rejected", "Omar Accountant"),
        ("Submitted", "Pending", "Pending Approval", "Field Engineer"),
    ]
    assert history[2]["reason"] == "After photos are missing."
    assert all(row["created_at"] and row["project_id"] == project["id"] for row in history)
    assert "Gulf Towers" in history[0]["reference"] and "Cooling Tower" in history[0]["reference"] and "SO 777" in history[0]["reference"]

    other = new_project(coordinator, name="Unassigned Project")
    upload_completion_images(coordinator, other["id"])
    coordinator.patch(f"/api/projects/{other['id']}", headers=csrf(coordinator), json={
        "completion_date": "2026-11-10",
        "completion_engineer_name": "Engineer", "invoice_number": "INV-OTHER",
    })
    assert {row["project_id"] for row in employee.get("/api/approval-history").json["data"]} == {project["id"]}
    assert len(coordinator.get("/api/approval-history").json["data"]) == 5

    assert coordinator.delete(f"{url}/permanent", headers=csrf(coordinator)).status_code == 200
    kept = [row for row in coordinator.get("/api/approval-history").json["data"] if "SO 777" in row["reference"]]
    assert len(kept) == 4 and all(row["project_id"] is None for row in kept)


def test_project_map_locations_are_saved_validated_and_scoped(app, client):
    headers = csrf(client)
    located = new_project(client, name="Located Site", latitude=25.204849, longitude=55.270783)
    unlocated = new_project(client, name="No Location Site")
    for invalid in ({"latitude": 95, "longitude": 10}, {"latitude": 25.2}, {"latitude": "25.2", "longitude": "55.2"},
                    {"latitude": 25.2, "longitude": None}):
        assert client.patch(f"/api/projects/{unlocated['id']}", headers=headers, json=invalid).status_code == 400, invalid

    pins = client.get("/api/dashboard/locations").json["data"]
    assert [pin["id"] for pin in pins] == [located["id"]]
    pin = pins[0]
    assert (pin["latitude"], pin["longitude"]) == (25.204849, 55.270783)
    assert (pin["client_name"], pin["name"], pin["location"], pin["status"]) == ("Gulf Towers", "Located Site", "Tower A", "Planned")
    assert pin["engineer_name"] == "Unassigned"

    assert client.patch(f"/api/projects/{unlocated['id']}", headers=headers, json={"latitude": 24.4539, "longitude": 54.3773}).status_code == 200
    assert len(client.get("/api/dashboard/locations").json["data"]) == 2
    cleared = client.patch(f"/api/projects/{located['id']}", headers=headers, json={"latitude": None, "longitude": None})
    assert cleared.status_code == 200 and cleared.json["data"]["latitude"] is None
    assert [pin["id"] for pin in client.get("/api/dashboard/locations").json["data"]] == [unlocated["id"]]

    employee, profile = employee_with_engineer(app, client, "Map Engineer", "ENG-020")
    assert employee.get("/api/dashboard/locations").json["data"] == []
    client.patch(f"/api/projects/{unlocated['id']}", headers=headers, json={"assigned_engineer_ids": [profile["id"]]})
    assert [pin["id"] for pin in employee.get("/api/dashboard/locations").json["data"]] == [unlocated["id"]]


def migration_app(database_path):
    class MigrationConfig:
        TESTING = True
        SECRET_KEY = "migration-test"
        SQLALCHEMY_DATABASE_URI = f"sqlite:///{database_path.as_posix()}"
        STORAGE_PATH = str(database_path.parent / "uploads")

    return create_app(MigrationConfig)


def test_migration_adds_new_fields_and_keeps_existing_records(tmp_path):
    from flask_migrate import downgrade, upgrade
    from sqlalchemy import inspect, text

    application = migration_app(tmp_path / "upgrade.db")
    with application.app_context():
        upgrade()
        downgrade(revision="202610240005")
        assert "approval_history" not in inspect(db.engine).get_table_names()
        now = datetime.now(timezone.utc).isoformat()
        with db.engine.begin() as connection:
            connection.execute(text("INSERT INTO organizations (id, name, created_at, updated_at) VALUES ('o', 'Org', :now, :now)"), {"now": now})
            connection.execute(text("INSERT INTO users (id, email, password_hash, is_active, created_at, updated_at) VALUES ('m', 'm@example.test', 'x', 1, :now, :now)"), {"now": now})
            connection.execute(text("INSERT INTO organization_memberships (id, organization_id, user_id, role, status, created_at) VALUES ('mm', 'o', 'm', 'manager', 'active', :now)"), {"now": now})
            connection.execute(text("INSERT INTO clients (id, organization_id, name, is_active, created_at, updated_at) VALUES ('c', 'o', 'Existing Client', 1, :now, :now)"), {"now": now})
            for project_id, status in (("p-open", "Planned"), ("p-done", "Completed")):
                connection.execute(text(
                    "INSERT INTO projects (id, organization_id, client_id, name, status, completion_status, engineer_visit, microbiology_frequency, created_at, updated_at) "
                    "VALUES (:id, 'o', 'c', :id, :status, :completion_status, 'Monthly', 'Weekly', :now, :now)"
                ), {
                    "id": project_id, "status": status,
                    "completion_status": "Completed Early" if status == "Completed" else None,
                    "now": now,
                })
            connection.execute(text("INSERT INTO engineers (id, organization_id, full_name, is_active, created_at, updated_at) VALUES ('e', 'o', 'Existing Engineer', 1, :now, :now)"), {"now": now})
        upgrade()
        inspector = inspect(db.engine)
        assert {
            "latitude", "longitude", "invoice_number", "approval_status", "rejection_reason",
            "engineer_visit_custom_date", "priority_date", "priority_rank",
        } <= {column["name"] for column in inspector.get_columns("projects")}
        assert "engineer_code" in {column["name"] for column in inspector.get_columns("engineers")}
        assert "approval_history" in inspector.get_table_names()
        with db.engine.connect() as connection:
            assert connection.execute(text("SELECT version_num FROM alembic_version")).scalar() == "202610300010"
            assert connection.execute(text("SELECT role FROM organization_memberships WHERE id = 'mm'")).scalar() == "service_coordinator"
            assert {"completion_date", "completion_engineer_name"} <= {column["name"] for column in inspector.get_columns("projects")}
            rows = dict(connection.execute(text("SELECT id, approval_status FROM projects")).all())
            assert rows == {"p-open": "Pending", "p-done": None}
            assert connection.execute(text("SELECT completion_status FROM projects WHERE id = 'p-done'")).scalar() == "Completed Before the Scheduled Date"
            assert connection.execute(text("SELECT engineer_visit || microbiology_frequency FROM projects WHERE id = 'p-open'")).scalar() == "MonthlyWeekly"
            assert connection.execute(text("SELECT count(*) FROM clients")).scalar() == 1
            assert connection.execute(text("SELECT full_name FROM engineers")).scalar() == "Existing Engineer"
        db.engine.dispose()


def test_records_persist_after_the_application_restarts(tmp_path):
    from flask_migrate import upgrade

    database_path = tmp_path / "persist.db"
    first = migration_app(database_path)
    with first.app_context():
        upgrade()
        now = datetime.now(timezone.utc).isoformat()
        db.session.execute(organizations.insert().values(id="org-p", name="Persist", created_at=now, updated_at=now))
        db.session.execute(users.insert().values(
            id="admin-p", email="persist@example.test", full_name="Persist Admin",
            password_hash=generate_password_hash(ROLE_PASSWORD), access_key_hash=generate_password_hash(ROLE_KEY),
            is_active=True, created_at=now, updated_at=now,
        ))
        db.session.execute(memberships.insert().values(
            id=str(uuid4()), organization_id="org-p", user_id="admin-p", role="administrator", status="active", created_at=now,
        ))
        db.session.commit()
        admin = sign_in(first, "persist@example.test", key=ROLE_KEY)
        engineer = admin.post("/api/engineers", headers=csrf(admin), json={"full_name": "Persisted Engineer", "engineer_code": "ENG-900"}).json["data"]
        project = new_project(admin, engineer_name="Persisted Engineer", assigned_engineer_ids=[engineer["id"]],
                              engineer_visit="Custom", engineer_visit_custom_date="2026-12-24", latitude=21.4858, longitude=39.1925)
        upload_completion_images(admin, project["id"])
        admin.patch(f"/api/projects/{project['id']}", headers=csrf(admin), json={
            "completion_date": "2026-11-10",
            "completion_engineer_name": "Persisted Engineer", "invoice_number": "INV-PERSIST",
        })
        admin.post(f"/api/projects/{project['id']}/reject", headers=csrf(admin), json={"reason": "Check invoice"})
        db.session.remove()
        db.engine.dispose()

    restarted = migration_app(database_path)
    with restarted.app_context():
        admin = sign_in(restarted, "persist@example.test", key=ROLE_KEY)
        saved = admin.get(f"/api/projects/{project['id']}").json["data"]
        assert (saved["engineer_visit_custom_date"], saved["invoice_number"], saved["approval_status"]) == ("2026-12-24", "INV-PERSIST", "Rejected")
        assert (saved["completion_date"], saved["completion_engineer_name"]) == ("2026-11-10", "Persisted Engineer")
        assert (saved["latitude"], saved["longitude"], saved["rejection_reason"]) == (21.4858, 39.1925, "Check invoice")
        assert admin.get("/api/engineers").json["data"][0]["engineer_code"] == "ENG-900"
        assert [row["action"] for row in admin.get("/api/approval-history").json["data"]] == ["Rejected", "Submitted"]
        db.engine.dispose()


# --- Work order (date order with manual moves) and list filters ---

def project_names(test_client, query=""):
    response = test_client.get(f"/api/projects?{query}")
    assert response.status_code == 200, response.json
    return [row["name"] for row in response.json["data"]]


def test_projects_list_in_date_order_and_can_be_moved_anywhere(app, client):
    headers = csrf(client)
    early = new_project(client, name="Early", start_date="2026-11-01", end_date="2026-11-01")
    middle = new_project(client, name="Middle", start_date="2026-11-03", end_date="2026-11-04")
    late = new_project(client, name="Late", start_date="2026-11-05", end_date="2026-11-05")
    assert project_names(client) == ["Early", "Middle", "Late"]

    moved = client.post(f"/api/projects/{late['id']}/move", headers=headers, json={"above_id": None, "below_id": early["id"]})
    assert moved.status_code == 200, moved.json
    assert moved.json["data"]["start_date"] == "2026-11-05"
    assert project_names(client) == ["Late", "Early", "Middle"]
    assert client.post(f"/api/projects/{middle['id']}/move", headers=headers, json={"above_id": late["id"], "below_id": early["id"]}).status_code == 200
    assert project_names(client) == ["Late", "Middle", "Early"]
    assert client.post(f"/api/projects/{late['id']}/move", headers=headers, json={"above_id": early["id"]}).status_code == 200
    assert project_names(client) == ["Middle", "Early", "Late"]

    new_project(client, name="New", start_date="2026-11-02", end_date="2026-11-02")
    assert project_names(client) == ["Middle", "Early", "Late", "New"]
    rescheduled = client.patch(f"/api/projects/{late['id']}", headers=headers, json={"start_date": "2026-11-10", "end_date": "2026-11-10"})
    assert rescheduled.status_code == 200 and rescheduled.json["data"]["priority_date"] is None
    assert project_names(client) == ["Middle", "Early", "New", "Late"]

    for invalid, status in (
        ({}, 400), ({"above_id": late["id"], "below_id": late["id"]}, 400), ({"below_id": middle["id"]}, 400),
        ({"below_id": "not-an-id"}, 400), ({"below_id": str(uuid4())}, 404),
    ):
        response = client.post(f"/api/projects/{middle['id']}/move", headers=headers, json=invalid)
        assert response.status_code == status, (invalid, response.json)

    for role in ("employee", "viewer"):
        user = add_account(app, client, role, f"Order {role}")
        assert user.post(f"/api/projects/{late['id']}/move", headers=csrf(user), json={"below_id": early["id"]}).status_code == 403
        assert user.post("/api/projects/order/reset", headers=csrf(user)).status_code == 403

    assert client.post("/api/projects/order/reset", headers=headers).status_code == 200
    assert project_names(client) == ["Early", "New", "Middle", "Late"]


def test_lists_filter_by_search_client_project_and_date(client):
    headers = csrf(client)
    pool = new_project(client, client_name="Alpha Hotels", name="Pool System", so_no=4401,
                       start_date="2026-11-01", end_date="2026-11-10", engineer_name="Rami Haddad")
    chiller = new_project(client, client_name="Beta Towers", name="Chiller Plant", so_no=5502,
                          start_date="2026-12-01", end_date="2026-12-02", location="Lusail Marina")
    assert project_names(client, "search=chiller") == ["Chiller Plant"]
    assert project_names(client, "search=alpha") == ["Pool System"]
    assert project_names(client, "search=4401") == ["Pool System"]
    assert project_names(client, "search=haddad") == ["Pool System"]
    assert project_names(client, "search=lusail") == ["Chiller Plant"]
    assert project_names(client, "search=%25") == []
    assert project_names(client, f"client_id={chiller['client_id']}") == ["Chiller Plant"]
    assert project_names(client, f"project_id={pool['id']}") == ["Pool System"]
    assert project_names(client, "from=2026-11-05&to=2026-11-06") == ["Pool System"]
    services = {row["name"]: row["service"] for row in client.get("/api/projects").json["data"]}
    kind = services["Pool System"]
    other = "Tank cleaning" if kind == "Water testing" else "Water testing"
    assert "Pool System" in project_names(client, f"service={kind}")
    assert all(services[name] == kind for name in project_names(client, f"service={kind}"))
    assert "Pool System" not in project_names(client, f"service={other}")
    assert client.get("/api/projects?service=Other").json["data"] == []
    assert project_names(client, "from=2026-11-11&to=2026-11-30") == []
    assert project_names(client, "from=2026-12-02") == ["Chiller Plant"]
    for bad in ("from=2026-13-01", "from=2026-11-10&to=2026-11-01", "client_id=nope", f"search={'x' * 101}"):
        assert client.get(f"/api/projects?{bad}").status_code == 400, bad

    cleaning = client.post("/api/cleaning-records", headers=headers, json={
        "client_id": pool["client_id"], "cleaning_date": "2026-11-03", "tank_identification": "T-9",
    }).json["data"]
    client.post("/api/cleaning-records", headers=headers, json={
        "client_id": chiller["client_id"], "cleaning_date": "2026-12-03", "tank_identification": "T-1",
    })
    cleaning_ids = lambda query: [row["id"] for row in client.get(f"/api/cleaning-records?{query}").json["data"]]
    assert cleaning_ids("search=alpha") == [cleaning["id"]]
    assert cleaning_ids("from=2026-11-01&to=2026-11-30") == [cleaning["id"]]
    assert client.get(f"/api/cleaning-records?project_id={pool['id']}").status_code == 400

    client.post("/api/engineers", headers=headers, json={"full_name": "Search Engineer", "engineer_code": "ENG-77"})
    assert [row["engineer_code"] for row in client.get("/api/engineers?search=eng-77").json["data"]] == ["ENG-77"]
    assert [row["name"] for row in client.get("/api/clients?search=towers").json["data"]] == ["Beta Towers"]
    assert client.get("/api/clients?from=2026-11-01").status_code == 400

    export = client.get("/api/reports/export?resource=projects&format=csv&search=chiller")
    assert export.status_code == 200
    assert b"Chiller Plant" in export.data and b"Pool System" not in export.data


# --- Admin login with a private key, client login, and employee logins made by administrators ---

def login_attempt(app, **body):
    return app.test_client().post("/api/auth/login", json=body)


def test_admin_login_needs_the_private_key_and_client_login_needs_only_email_and_password(app, client):
    admin = {"email": "admin@example.test", "password": "secure-test-password"}
    assert login_attempt(app, **admin, login_type="admin", access_key=ADMIN_KEY).status_code == 200
    assert login_attempt(app, **admin, login_type="admin").status_code == 400
    assert login_attempt(app, **admin, login_type="admin", access_key="wrong-private-key").status_code == 401
    assert login_attempt(app, email="admin@example.test", password="wrong", login_type="admin", access_key=ADMIN_KEY).status_code == 401
    refused = login_attempt(app, **admin)
    assert refused.status_code == 403 and "Admin login" in refused.json["error"]["message"]
    assert login_attempt(app, **admin, login_type="other").status_code == 400

    coordinator = add_account(app, client, "service_coordinator", "Key Coordinator")
    email = coordinator.get("/api/auth/me").json["user"]["email"]
    assert login_attempt(app, email=email, password=ROLE_PASSWORD).status_code == 403
    assert login_attempt(app, email=email, password=ROLE_PASSWORD, login_type="admin", access_key=ADMIN_KEY).status_code == 401
    assert login_attempt(app, email=email, password=ROLE_PASSWORD, login_type="admin", access_key=ROLE_KEY).status_code == 200

    add_account(app, client, "employee", "Plain Employee")
    employee_email = next(u["email"] for u in [client.get("/api/engineers").json["data"][0]] if u["email"])
    employee = login_attempt(app, email=employee_email, password=ROLE_PASSWORD)
    assert employee.status_code == 200 and employee.json["role"] == "employee"
    refused = login_attempt(app, email=employee_email, password=ROLE_PASSWORD, login_type="admin", access_key=ROLE_KEY)
    assert refused.status_code == 403 and "Client login" in refused.json["error"]["message"]

    headers = csrf(client)
    assert client.post("/api/auth/users", headers=headers, json={
        "email": "nokey@example.test", "full_name": "No Key", "password": ROLE_PASSWORD, "role": "accountant",
    }).status_code == 400
    assert client.post("/api/auth/users", headers=headers, json={
        "email": "keyed@example.test", "full_name": "Keyed", "password": ROLE_PASSWORD, "role": "employee", "access_key": ROLE_KEY,
    }).status_code == 400
    with app.app_context():
        stored = db.session.execute(db.select(users.c.access_key_hash).where(users.c.email == "admin@example.test")).scalar()
    assert stored and ADMIN_KEY not in stored


def test_administrator_creates_and_resets_employee_logins_from_the_engineers_page(app, client):
    headers = csrf(client)
    created = client.post("/api/engineers", headers=headers, json={
        "full_name": "Nadia Karim", "engineer_code": "ENG-300", "email": "Nadia@Gmail.test", "password": "nadia-login-pass-1",
    })
    assert created.status_code == 201, created.json
    assert created.json["data"]["user_id"] and "password" not in created.json["data"]
    signed_in = login_attempt(app, email="nadia@gmail.test", password="nadia-login-pass-1")
    assert signed_in.status_code == 200 and signed_in.json["role"] == "employee"

    assert client.post("/api/engineers", headers=headers, json={
        "full_name": "Copy", "engineer_code": "ENG-301", "email": "nadia@gmail.test", "password": "another-long-pass-1",
    }).status_code == 409
    assert client.post("/api/engineers", headers=headers, json={
        "full_name": "Short", "engineer_code": "ENG-302", "email": "short@gmail.test", "password": "short",
    }).status_code == 400
    assert client.post("/api/engineers", headers=headers, json={
        "full_name": "No Mail", "engineer_code": "ENG-303", "password": "a-long-enough-pass-1",
    }).status_code == 400

    reset = client.patch(f"/api/engineers/{created.json['data']['id']}", headers=headers, json={"password": "nadia-new-password-2"})
    assert reset.status_code == 200, reset.json
    assert login_attempt(app, email="nadia@gmail.test", password="nadia-login-pass-1").status_code == 401
    assert login_attempt(app, email="nadia@gmail.test", password="nadia-new-password-2").status_code == 200

    plain = client.post("/api/engineers", headers=headers, json={"full_name": "Later Login", "engineer_code": "ENG-304", "email": "later@gmail.test"}).json["data"]
    assert plain["user_id"] is None
    assert client.patch(f"/api/engineers/{plain['id']}", headers=headers, json={"password": "later-login-pass-1"}).status_code == 200
    assert login_attempt(app, email="later@gmail.test", password="later-login-pass-1").status_code == 200

    coordinator = add_account(app, client, "service_coordinator", "Not Admin")
    assert coordinator.post("/api/engineers", headers=csrf(coordinator), json={
        "full_name": "Blocked", "engineer_code": "ENG-305", "email": "blocked@gmail.test", "password": "blocked-login-pass-1",
    }).status_code == 403
    assert coordinator.post("/api/engineers", headers=csrf(coordinator), json={"full_name": "Fine", "engineer_code": "ENG-306"}).status_code == 201


def test_engineers_upload_profile_photos_with_the_right_permissions(app, client):
    headers = csrf(client)
    created = client.post("/api/engineers", headers=headers, json={
        "full_name": "Photo Person", "engineer_code": "ENG-400", "email": "photo@gmail.test", "password": "photo-login-pass-1",
    }).json["data"]
    other = client.post("/api/engineers", headers=headers, json={"full_name": "Someone Else", "engineer_code": "ENG-401"}).json["data"]
    png = b"\x89PNG\r\n\x1a\nprofile"

    def send(test_client, engineer_id, data=png, name="me.png", image_type="profile"):
        return test_client.post("/api/uploads", headers=csrf(test_client), content_type="multipart/form-data", data={
            "related_type": "engineer", "related_id": engineer_id, "image_type": image_type, "file": (BytesIO(data), name, "image/png"),
        })

    employee = sign_in(app, "photo@gmail.test", "photo-login-pass-1")
    first = send(employee, created["id"])
    assert first.status_code == 201, first.json
    assert send(employee, other["id"]).status_code in {403, 404}
    assert send(employee, created["id"], b"not an image", "me.png").status_code == 400
    assert send(employee, created["id"], b"%PDF-1.4 x", "me.pdf").status_code == 400
    assert send(employee, created["id"], image_type="before").status_code == 400

    listed = {row["id"]: row["photo_id"] for row in client.get("/api/engineers").json["data"]}
    assert listed[created["id"]] == first.json["data"]["id"] and listed[other["id"]] is None
    assert employee.get(f"/api/uploads/{first.json['data']['id']}?inline=1").status_code == 200

    second = send(employee, created["id"], png + b"2", "new.png")
    assert second.status_code == 201
    assert client.get(f"/api/uploads/{first.json['data']['id']}").status_code == 404
    assert {row["id"]: row["photo_id"] for row in client.get("/api/engineers").json["data"]}[created["id"]] == second.json["data"]["id"]

    assert send(client, other["id"]).status_code == 201
    viewer = add_account(app, client, "viewer", "Photo Viewer")
    assert send(viewer, other["id"]).status_code == 403
    assert viewer.post("/api/uploads", headers=csrf(viewer), content_type="multipart/form-data", data={
        "related_type": "client", "related_id": str(uuid4()), "file": (BytesIO(b"%PDF-1.4 x"), "a.pdf", "application/pdf"),
    }).status_code == 403


def test_projects_export_uses_the_client_schedule_layout(client):
    from io import BytesIO
    from openpyxl import load_workbook
    new_project(client, client_name="Daikin", name="True Bell", so_no=1087, start_date="2026-07-01", end_date="2027-03-31",
                system_details="CT - 2 samples\nCHW 1 - 1 samples", engineer_visit="Quarterly",
                microbiology="Yes", microbiology_frequency="Quarterly", legionella="No")
    new_project(client, client_name="Daikin", name="Chef Middle East", so_no=846, start_date="2026-01-01", end_date="2026-12-31")
    new_project(client, client_name="Acme", name="Solo", so_no=7, start_date="2026-05-01", end_date="2026-05-20")
    response = client.get("/api/reports/export?resource=projects&format=xlsx")
    assert response.status_code == 200
    sheet = load_workbook(BytesIO(response.data)).active
    assert [cell.value for cell in sheet[1]] == [
        "SL NO", "Client", None, "Project", "SO Number", "Contract Period", "System Details",
        "Engineer Visit", "Microbiology", "Legionella",
    ]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]
    assert rows[0][:2] == [1, "Acme"] and rows[0][5] == "May 2026"
    assert rows[1][:5] == [2, "Daikin", 1, "Chef Middle East", 846]
    assert rows[2][:5] == [None, None, 2, "True Bell", 1087]
    assert rows[2][5] == "July 2026 to March 2027"
    assert rows[2][6] == "CT - 2 samples\nCHW 1 - 1 samples"
    assert rows[2][7:] == ["Quarterly", "Yes quarterly", "NO"]
    assert {str(r) for r in sheet.merged_cells.ranges} == {"A3:A4", "B3:B4"}
    assert sheet["B2"].fill.fgColor.rgb.endswith("FFFF00") and sheet["B2"].font.bold
    for fmt in ("csv", "pdf"):
        assert client.get(f"/api/reports/export?resource=projects&format={fmt}").status_code == 200


def _certificate_file(text_runs):
    """A one-slide deck whose text is split across several runs, the way PowerPoint often stores it."""
    from io import BytesIO
    from pptx import Presentation
    from pptx.util import Inches
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(2)).text_frame
    for index, runs in enumerate(text_runs):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        for text in runs:
            paragraph.add_run().text = text
        paragraph.runs[0].font.bold = True
    out = BytesIO()
    deck.save(out)
    return out.getvalue()


def upload_template(test_client, service, deck, name="Certificate", filename="cert.pptx"):
    from io import BytesIO
    return test_client.post("/api/certificate-templates", headers=csrf(test_client), content_type="multipart/form-data",
                            data={"name": name, "service": service, "file": (BytesIO(deck), filename)})


def test_services_page_defaults_create_rename_and_in_use_rules(client):
    headers = csrf(client)
    assert [row["name"] for row in client.get("/api/services").json["data"]] == ["Tank cleaning", "Water testing"]

    created = client.post("/api/services", headers=headers, json={"name": "Legionella audit", "certificate_text": "Cooling Tower System", "validity_months": 3})
    assert created.status_code == 201, created.json
    service = created.json["data"]
    assert service["is_active"] is True
    assert client.post("/api/services", headers=headers, json={"name": "legionella AUDIT"}).status_code == 409
    assert client.post("/api/services", headers=headers, json={"name": "Bad", "validity_months": 0}).status_code == 400

    assert client.post("/api/projects", headers=headers, json={"client_name": "C", "name": "P", "service": "Not a service"}).status_code == 400
    project = new_project(client, service="Legionella audit")
    assert client.get(f"/api/projects?service=Legionella audit").json["data"][0]["id"] == project["id"]

    renamed = client.patch(f"/api/services/{service['id']}", headers=headers, json={"name": "Legionella inspection"})
    assert renamed.status_code == 200
    assert client.get(f"/api/projects/{project['id']}").json["data"]["service"] == "Legionella inspection"
    assert client.delete(f"/api/services/{service['id']}", headers=headers).status_code == 409

    client.patch(f"/api/services/{service['id']}", headers=headers, json={"is_active": False})
    assert client.post("/api/projects", headers=headers, json={"client_name": "C", "name": "P2", "service": "Legionella inspection"}).status_code == 400
    unused = client.post("/api/services", headers=headers, json={"name": "Spare"}).json["data"]
    assert client.delete(f"/api/services/{unused['id']}", headers=headers).status_code == 200


def test_services_are_administrator_only_to_change(app, client):
    coordinator = add_account(app, client, "service_coordinator", "Sara Coordinator")
    viewer = add_account(app, client, "viewer", "Vic Viewer")
    assert viewer.get("/api/services").status_code == 200
    assert coordinator.post("/api/services", headers=csrf(coordinator), json={"name": "Nope"}).status_code == 403
    assert viewer.post("/api/services", headers=csrf(viewer), json={"name": "Nope"}).status_code == 403


def test_certificate_is_generated_for_the_client_after_admin_approval(app, client, monkeypatch):
    from io import BytesIO
    from pptx import Presentation
    converted = {}

    def fake_pdf(data, extension):
        converted["data"], converted["extension"] = data, extension
        return b"%PDF-1.4 certificate"

    monkeypatch.setattr("backend.routes.api.convert_to_pdf", fake_pdf)
    headers = csrf(client)
    deck = _certificate_file([["{{client", "_name}}"], ["{{project_name}} | {{location}}"], ["{{service_text}}"], ["{{completion_date}} - {{expiry_date}}"]])

    assert upload_template(client, "", deck).status_code == 400
    assert upload_template(client, "Unknown service", deck).status_code == 400
    assert upload_template(client, "Tank cleaning", b"not a deck").status_code == 400
    assert upload_template(client, "Tank cleaning", b"hello", filename="note.txt").status_code == 400
    saved = upload_template(client, "Tank cleaning", deck, name="Disinfection")
    assert saved.status_code == 201, saved.json
    template = saved.json["data"]
    assert template["service"] == "Tank cleaning" and "tags" not in template and "storage_path" not in template
    assert [(row["name"], row["service"]) for row in client.get("/api/certificate-templates").json["data"]] == [("Disinfection", "Tank cleaning")]

    employee, profile = employee_with_engineer(app, client, "Field Engineer", "ENG-077")
    project = new_project(client, client_name="IMDAAD LLC", name="VENETO VILLA - 01", service="Tank cleaning",
                          location="Veneto Villa, Dubai, UAE", engineer_name="Field Engineer", assigned_engineer_ids=[profile["id"]])
    url = f"/api/projects/{project['id']}"
    assert employee.get(f"{url}/certificate").status_code == 409

    upload_completion_images(employee, project["id"])
    assert employee.patch(url, headers=csrf(employee), json={
        "completion_date": "2026-09-09", "completion_engineer_name": "Field Engineer", "invoice_number": "INV-9",
    }).status_code == 200
    assert employee.get(f"{url}/certificate").status_code == 409
    assert client.post(f"{url}/approve", headers=headers).status_code == 200

    certificate = employee.get(f"{url}/certificate")
    assert certificate.status_code == 200
    assert certificate.mimetype == "application/pdf" and certificate.data.startswith(b"%PDF")
    assert "certificate-VENETO_VILLA_-_01.pdf" in certificate.headers["Content-Disposition"]
    filled = Presentation(BytesIO(converted["data"])).slides[0]
    assert converted["extension"] == "pptx"
    assert all(run.font.bold is False for shape in filled.shapes for p in shape.text_frame.paragraphs for run in p.runs)
    text = "\n".join(p.text for shape in filled.shapes for p in shape.text_frame.paragraphs)
    assert "IMDAAD LLC" in text and "VENETO VILLA - 01 | Veneto Villa, Dubai, UAE" in text
    assert "Domestic Water Tank and Pipeline Network." in text
    assert "09/09/2026 - 09/03/2027" in text and "{{" not in text

    def no_converter(data, extension):
        from backend.utils.pdf_convert import PdfUnavailable
        raise PdfUnavailable("none")

    monkeypatch.setattr("backend.routes.api.convert_to_pdf", no_converter)
    assert employee.get(f"{url}/certificate").status_code == 503
    assert client.get(f"/api/certificate-templates/{template['id']}/preview").status_code == 503
    monkeypatch.setattr("backend.routes.api.convert_to_pdf", fake_pdf)
    preview = client.get(f"/api/certificate-templates/{template['id']}/preview")
    assert preview.status_code == 200 and preview.mimetype == "application/pdf"
    assert "Sample Client LLC" in "\n".join(p.text for shape in Presentation(BytesIO(converted["data"])).slides[0].shapes for p in shape.text_frame.paragraphs)
    assert employee.get(f"/api/certificate-templates/{template['id']}/preview").status_code == 403
    other = new_project(client, service="Water testing")
    assert client.get(f"/api/projects/{other['id']}/certificate").status_code == 409
    assert client.delete(f"/api/certificate-templates/{template['id']}", headers=headers).status_code == 200
    assert client.get(f"{url}/certificate").status_code == 404
    assert client.get("/api/certificate-templates").json["data"] == []


class FakeR2:
    """Stands in for the R2 binding (env.FILES): keeps objects in memory, like the Workers SDK wrapper returns them."""

    def __init__(self):
        self.objects = {}

    def put(self, key, data, options=None):
        self.objects[key] = bytes(data)

    def get(self, key):
        if key not in self.objects:
            return None
        body = self.objects[key]

        class Stored:
            def arrayBuffer(self):
                return memoryview(body)

        return Stored()

    def delete(self, key):
        self.objects.pop(key, None)


def pretend_worker(monkeypatch, env):
    """Run the Worker-only code paths in CPython: fake the small part of Pyodide (`js`, `pyodide.ffi`) they use."""
    import sys
    import types
    js = types.ModuleType("js")
    js.Object = types.SimpleNamespace(fromEntries=dict)
    js.JSON = types.SimpleNamespace(parse=lambda text: None)
    ffi = types.ModuleType("pyodide.ffi")
    ffi.run_sync = lambda value: value
    ffi.to_js = lambda value, **kwargs: value
    pyodide = types.ModuleType("pyodide")
    pyodide.ffi = ffi
    monkeypatch.setitem(sys.modules, "js", js)
    monkeypatch.setitem(sys.modules, "pyodide", pyodide)
    monkeypatch.setitem(sys.modules, "pyodide.ffi", ffi)
    monkeypatch.setattr("backend.utils.storage.in_workers", lambda: True)
    monkeypatch.setattr("backend.utils.storage.worker_env", lambda: env)


def test_uploads_and_certificate_templates_are_kept_in_the_r2_bucket(app, client, monkeypatch):
    from types import SimpleNamespace
    bucket = FakeR2()
    project = new_project(client, name="Cloud Project")
    pretend_worker(monkeypatch, SimpleNamespace(FILES=bucket))
    png = b"\x89PNG\r\n\x1a\n" + b"0" * 20
    upload = client.post("/api/uploads", headers=csrf(client), content_type="multipart/form-data", data={
        "related_type": "project", "related_id": project["id"], "image_type": "before", "file": (BytesIO(png), "site.png"),
    })
    assert upload.status_code == 201, upload.json
    stored = upload.json["data"]["storage_path"]
    assert bucket.objects[stored] == png
    assert not (Path(app.config["STORAGE_PATH"]) / stored).exists()

    assert client.get(f"/api/uploads/{upload.json['data']['id']}").data == png
    assert client.delete(f"/api/uploads/{upload.json['data']['id']}", headers=csrf(client)).status_code == 200
    assert stored not in bucket.objects

    monkeypatch.setattr("backend.routes.api.convert_to_pdf", lambda data, extension: b"%PDF-1.4")
    deck = _certificate_file([["{{client_name}}"]])
    template = upload_template(client, "Water testing", deck).json["data"]
    assert deck in bucket.objects.values()
    assert client.get(f"/api/certificate-templates/{template['id']}/preview").status_code == 200
    assert client.delete(f"/api/certificate-templates/{template['id']}", headers=csrf(client)).status_code == 200
    assert deck not in bucket.objects.values()


def test_r2_keys_reject_path_tricks(app, monkeypatch):
    from types import SimpleNamespace
    from backend.utils import storage
    pretend_worker(monkeypatch, SimpleNamespace(FILES=FakeR2()))
    with app.app_context():
        for bad in ("../x", "a//b", "/abs", "a/./b", ""):
            with pytest.raises(storage.StorageError):
                storage.save(bad, b"x", "text/plain")


def test_worker_passwords_use_web_crypto_pbkdf2_in_werkzeug_format(monkeypatch):
    import hashlib
    from werkzeug.security import check_password_hash, generate_password_hash
    from backend.security import passwords
    monkeypatch.setattr(passwords, "in_workers", lambda: True)
    monkeypatch.setattr(passwords, "_webcrypto_pbkdf2", lambda password, salt, iterations, hash_name="SHA-256", length=32: hashlib.pbkdf2_hmac(
        hash_name.replace("-", "").lower(), password.encode(), salt.encode(), iterations, length).hex())
    stored = passwords.hash_password("worker-password-1")
    assert stored.startswith("pbkdf2:sha256:600000$")
    assert passwords.verify_password(stored, "worker-password-1")
    assert not passwords.verify_password(stored, "wrong-password-1")
    assert check_password_hash(stored, "worker-password-1")  # the same hash still works on a normal server
    assert passwords.verify_password(generate_password_hash("local-made", method="pbkdf2:sha256:1000"), "local-made")
    assert not passwords.verify_password(generate_password_hash("old-scrypt", method="scrypt"), "old-scrypt")


def test_worker_config_uses_the_d1_driver():
    from config import WorkerConfig
    assert WorkerConfig.SQLALCHEMY_DATABASE_URI == "sqlite+d1://"
    assert WorkerConfig.SESSION_COOKIE_SECURE is True and WorkerConfig.SECRET_KEY is None


def test_d1_driver_reports_changes_and_returns_rows():
    from sqlalchemy import create_engine, text
    from backend import d1
    d1.use_runner(d1.SQLiteRunner())
    try:
        engine = create_engine("sqlite+d1://")
        with engine.begin() as connection:
            connection.execute(text("CREATE TABLE t (id TEXT PRIMARY KEY, n INTEGER)"))
            assert connection.execute(text("INSERT INTO t VALUES ('a', 1), ('b', 2)")).rowcount == 2
            assert connection.execute(text("UPDATE t SET n = n + 1 WHERE id = :id"), {"id": "a"}).rowcount == 1
            assert connection.execute(text("SELECT id, n FROM t ORDER BY id")).mappings().all() == [{"id": "a", "n": 2}, {"id": "b", "n": 2}]
            assert connection.execute(text("DELETE FROM t WHERE id = 'b' RETURNING id")).scalar() == "b"
            assert connection.execute(text("SELECT id FROM t WHERE id = 'zzz'")).first() is None
            with pytest.raises(Exception):
                connection.execute(text("INSERT INTO t VALUES ('a', 9)"))
    finally:
        d1.use_runner(None)


def test_empty_database_reports_not_set_up_instead_of_a_server_error():
    if USE_D1:
        from backend import d1
        d1.use_runner(d1.SQLiteRunner())
    application = create_app(TestConfig)
    with application.app_context():
        test_client = application.test_client()
        assert test_client.get("/api/health").json == {"status": "database_not_set_up"}
        response = test_client.post("/api/auth/login", json={
            "email": "owner@example.test", "password": "secure-test-password", "login_type": "admin", "access_key": ADMIN_KEY,
        })
        assert response.status_code == 503
        assert "not set up" in response.json["error"]["message"]


def test_d1_migration_creates_every_table_and_column_the_app_uses():
    import sqlite3
    sql = (Path(__file__).parent.parent / "d1-migrations" / "0001_initial_schema.sql").read_text(encoding="utf-8")
    connection = sqlite3.connect(":memory:")
    connection.executescript(sql)
    for table in db.metadata.sorted_tables:
        columns = {row[1] for row in connection.execute(f"PRAGMA table_info({table.name})")}
        # Old columns kept by earlier migrations (projects.start_time, end_time) may remain; every current one must exist.
        assert {column.name for column in table.columns} <= columns, table.name
